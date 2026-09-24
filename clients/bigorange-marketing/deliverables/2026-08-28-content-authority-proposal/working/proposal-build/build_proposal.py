from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\dillo\Documents\Codex\projects\client-operations\clients\bigorange-marketing\deliverables\2026-08-28-content-authority-proposal")
OUT = ROOT / "outputs" / "01a04906-96f3-7072-b3bf-0c2589aed6fa" / "BigOrange-Custom-Home-Builder-Authority-System-Proposal-2026-08-28.docx"
PLAN_PATH = ROOT / "working" / "content-plan.json"

ORANGE = RGBColor(244, 119, 33)
ORANGE_DARK = RGBColor(183, 74, 8)
INK = RGBColor(17, 17, 17)
CHARCOAL = RGBColor(55, 55, 55)
GRAY = RGBColor(104, 104, 104)
LIGHT_GRAY = RGBColor(244, 246, 249)
PAPER = RGBColor(250, 247, 241)
PALE_ORANGE = RGBColor(255, 239, 224)
PALE_BLUE = RGBColor(224, 242, 249)
PALE_GREEN = RGBColor(231, 246, 236)
PALE_RED = RGBColor(252, 232, 230)
WHITE = RGBColor(255, 255, 255)

FONT = "Arial"
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def rgb_hex(color: RGBColor) -> str:
    return str(color)


def set_cell_fill(cell, color: RGBColor):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), rgb_hex(color))


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        border = borders.find(tag)
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa):
    assert sum(widths_dxa) == USABLE_DXA, (widths_dxa, sum(widths_dxa))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(USABLE_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        prevent_row_split(row)
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            set_cell_border(cell)


def set_run_font(run, size=None, color=None, bold=None, italic=None, all_caps=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if all_caps is not None:
        run.font.all_caps = all_caps


def set_paragraph_keep(paragraph, keep_next=False, keep_lines=True):
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    if keep_lines:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("PAGE ")
    set_run_font(run, size=8.5, color=GRAY, bold=True)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    r = paragraph.add_run()
    r._r.append(fld_char_1)
    r._r.append(instr)
    r._r.append(fld_char_2)


def shade_paragraph(paragraph, fill: RGBColor, border_color=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), rgb_hex(fill))
    if border_color:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), rgb_hex(border_color))
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_body(doc, text, bold_lead=None, italic=False, after=8, keep=False):
    p = doc.add_paragraph(style="Body")
    p.paragraph_format.space_after = Pt(after)
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, size=10.5, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.5, color=INK, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=INK, italic=italic)
    if keep:
        set_paragraph_keep(p)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(text, style=f"Heading {level}")
    set_paragraph_keep(p, keep_next=True)
    return p


def add_kicker(doc, text, centered=False):
    p = doc.add_paragraph(style="Kicker")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text.upper())
    set_run_font(r, size=8.5, color=ORANGE_DARK, bold=True, all_caps=True)
    set_paragraph_keep(p, keep_next=True)
    return p


def add_bullet(doc, text, level=0, bold_lead=None):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(" " + bold_lead)
        set_run_font(r1, size=10.25, color=INK, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2, size=10.25, color=INK)
    else:
        r = p.add_run(" " + text)
        set_run_font(r, size=10.25, color=INK)
    return p


def new_numbering_instance(doc, start=1):
    numbering = doc.part.numbering_part.element
    base_num_id = int(doc.styles["List Number"]._element.pPr.numPr.numId.val)
    base_num = next(
        node for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == base_num_id
    )
    abstract_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    used_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_id = max(used_ids) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_id


def add_numbered(doc, title, detail, num_id):
    p = doc.add_paragraph(style="List Number")
    p_pr = p._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    p_pr.append(num_pr)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.208
    r1 = p.add_run(f" {title}. ")
    set_run_font(r1, size=10.25, color=INK, bold=True)
    r2 = p.add_run(detail)
    set_run_font(r2, size=10.25, color=INK)
    return p


def add_callout(doc, label, text, tone="orange"):
    fills = {
        "orange": (PALE_ORANGE, ORANGE),
        "blue": (PALE_BLUE, RGBColor(43, 123, 155)),
        "green": (PALE_GREEN, RGBColor(38, 130, 72)),
        "red": (PALE_RED, RGBColor(178, 54, 43)),
        "gray": (LIGHT_GRAY, GRAY),
    }
    fill, accent = fills[tone]
    p = doc.add_paragraph(style="Callout")
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.right_indent = Inches(0.13)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    shade_paragraph(p, fill, accent)
    r1 = p.add_run(f"{label.upper()}  ")
    set_run_font(r1, size=9.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.25, color=INK)
    set_paragraph_keep(p)
    return p


def add_table(doc, headers, rows, widths_dxa, header_fill=ORANGE, banded=True, font_size=8.6):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_fill(cell, header_fill)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(str(header))
        set_run_font(r, size=font_size, color=WHITE if header_fill == ORANGE or header_fill == INK else INK, bold=True)
    for row_idx, row_data in enumerate(rows):
        cells = table.add_row().cells
        for idx, value in enumerate(row_data):
            if banded and row_idx % 2 == 0:
                set_cell_fill(cells[idx], PAPER)
            p = cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(value))
            set_run_font(r, size=font_size, color=INK)
    set_table_geometry(table, widths_dxa)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_metric_strip(doc, items):
    table = doc.add_table(rows=1, cols=len(items))
    table.style = "Table Grid"
    set_repeat_table_header(table.rows[0])
    widths = [USABLE_DXA // len(items)] * len(items)
    widths[-1] += USABLE_DXA - sum(widths)
    for idx, (value, label) in enumerate(items):
        cell = table.cell(0, idx)
        set_cell_fill(cell, INK if idx % 2 == 0 else ORANGE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(value)
        set_run_font(r, size=18, color=WHITE, bold=True)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(label.upper())
        set_run_font(r2, size=7.7, color=WHITE, bold=True, all_caps=True)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)
    return table


def add_section_break(doc, title, subtitle=None):
    kicker = add_kicker(doc, "BigOrange authority system")
    kicker.paragraph_format.page_break_before = True
    add_heading(doc, title, 1)
    if subtitle:
        p = doc.add_paragraph(style="Section Lead")
        r = p.add_run(subtitle)
        set_run_font(r, size=11.5, color=CHARCOAL, italic=True)
        p.paragraph_format.space_after = Pt(12)
    return p if subtitle else None


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333

    if "Body" not in styles:
        body = styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)
    else:
        body = styles["Body"]
    body.base_style = normal
    body.font.name = FONT
    body.font.size = Pt(10.5)
    body.font.color.rgb = INK
    body.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.space_after = Pt(8)
    body.paragraph_format.line_spacing = 1.333

    for level, size, before, after, color in (
        (1, 20, 18, 10, INK),
        (2, 14, 12, 6, ORANGE_DARK),
        (3, 11.5, 8, 4, CHARCOAL),
    ):
        st = styles[f"Heading {level}"]
        st.font.name = FONT
        st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.page_break_before = False

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        st = styles[name]
        st.font.name = FONT
        st.font.size = Pt(10.25)
        st.font.color.rgb = INK
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.208
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.375)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.194)
    styles["List Bullet 2"].paragraph_format.left_indent = Inches(0.65)
    styles["List Bullet 2"].paragraph_format.first_line_indent = Inches(-0.194)
    styles["List Number"].paragraph_format.left_indent = Inches(0.375)
    styles["List Number"].paragraph_format.first_line_indent = Inches(-0.194)

    for name in ("Kicker", "Callout", "Section Lead"):
        if name not in styles:
            st = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            st = styles[name]
        st.base_style = normal
        st.font.name = FONT


def configure_section(section, first=False):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82 if not first else 0.72)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.32)

    header = section.header
    p = header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("BIGORANGE.MARKETING  |  CUSTOM HOME BUILDER AUTHORITY SYSTEM")
    set_run_font(r, size=7.5, color=GRAY, bold=True, all_caps=True)
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), rgb_hex(ORANGE))
    p_bdr.append(bottom)
    p_pr.append(p_bdr)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.paragraph_format.space_after = Pt(0)
    left = fp.add_run("Leadership decision packet  |  August 28, 2026")
    set_run_font(left, size=8, color=GRAY)
    tab_stops = fp.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Inches(6.0))
    fp.add_run("\t")
    add_page_field(fp)


def word_count(path: Path) -> int:
    text = path.read_text(encoding="utf-8")
    text = re.sub(r"^---.*?---", "", text, count=1, flags=re.S)
    return len(re.findall(r"\b[\w'-]+\b", text))


def build():
    plan = json.loads(PLAN_PATH.read_text(encoding="utf-8"))
    assets = plan["assets"]
    current_assets = [a for a in assets if a["build_window"] == "Current pilot"]
    next_assets = [a for a in assets if "next phase" in a["build_window"].lower() or "days" in a["build_window"].lower()]
    reject_count = sum(1 for owner in plan["keyword_owner"].values() if owner.startswith("REJECT"))

    counts = {
        "HUB-01": word_count(ROOT / "content" / "pillar-page-interview-integrated.md"),
        "ART-01": word_count(ROOT / "content" / "supporting-article-01-interview-integrated.md"),
        "ART-02": word_count(ROOT / "content" / "supporting-article-02-interview-integrated.md"),
    }

    doc = Document()
    configure_styles(doc)
    configure_section(doc.sections[0], first=True)

    props = doc.core_properties
    props.title = "BigOrange Custom Home Builder Authority System"
    props.subject = "Final pilot proposal and leadership decision packet"
    props.author = "DM Marketing Specialist"
    props.keywords = "BigOrange, home builder marketing, SEO, AI search, content strategy, proposal"
    props.comments = "Prepared from the reconciled pilot, live-site verification, current keyword evidence, and Janice Weiser's August 28, 2026 interview."

    # Cover / proposal centerpiece
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BIGORANGE.MARKETING")
    set_run_font(r, size=11, color=ORANGE_DARK, bold=True, all_caps=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("Custom Home Builder\nAuthority System")
    set_run_font(r, size=29, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Final Pilot Proposal and Leadership Decision Packet")
    set_run_font(r, size=14, color=CHARCOAL, italic=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(26)
    r = p.add_run("Interview-integrated content, keyword ownership, release plan, measurement, and next-phase roadmap")
    set_run_font(r, size=9.75, color=GRAY, bold=True)

    add_metric_strip(doc, [
        ("62", "keyword decisions"),
        ("19", "owned assets"),
        ("3", "pilot drafts"),
        ("90", "day release plan"),
    ])

    add_callout(
        doc,
        "Decision purpose",
        "Approve the exact authority-hub direction and review gates, then release one coherent builder-marketing system without fragmenting the existing commercial URL.",
        "orange",
    )

    cover_rows = [
        ["Prepared for", "Margee, Paula, Janice, and BigOrange leadership"],
        ["Prepared by", "Dillon Mohr, DM Marketing Specialist"],
        ["Prepared", "August 28, 2026"],
        ["Pilot truth", "35 hours at $30/hour; approved total $1,050"],
        ["Decision date", "September 3, 2026 leadership review"],
        ["Status", "Complete proposal package; publication remains approval gated"],
    ]
    add_table(doc, ["Proposal field", "Confirmed position"], cover_rows, [2300, 7060], header_fill=INK, banded=False, font_size=9.4)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("CONFIDENTIAL WORKING PROPOSAL  |  NO PUBLIC CLAIM OR PUBLICATION IS IMPLIED")
    set_run_font(r, size=8, color=GRAY, bold=True, all_caps=True)

    # Executive decision
    add_section_break(doc, "Executive decision", "What leadership is approving and why it is ready now.")
    add_callout(
        doc,
        "Recommendation",
        "Approve one authority hub at the existing /marketing-agency-for-builders/ URL, select a single visual direction, and release the pillar plus two supporting articles only after Janice closes factual and public-use permissions.",
        "green",
    )
    add_heading(doc, "The system being approved", 2)
    add_body(doc, "This pilot converts a broad builder-services page into a focused authority system for custom home builders while preserving the existing URL that already owns relevant commercial intent. The system is deliberately smaller than the full keyword universe: one hub, distinct supporting reader jobs, explicit internal links, governed evidence, and a review loop that learns from real queries and qualified actions.")
    add_body(doc, "The final package now includes the audit and research foundation, the 62-query decision ledger, a 19-asset ownership map, the revised pillar, two complete supporting articles, a visible FAQ and schema candidate, a 90-day release sequence, a 12-month roadmap, AI-search implementation guidance, and publication QA gates.")

    add_heading(doc, "Why this is stronger after Janice's interview", 2)
    for text in (
        "The audience is framed around real builder growth moments: uneven pipeline, referrals slowing, an aging website, stretched internal marketing capacity, succession, or a shift into a new service mix.",
        "Trust is treated as a proof sequence, not a slogan: excellent project photography, current digital presence, customer evidence, a clear process, direct answers, and a low-pressure next step.",
        "The content answers questions builders commonly leave unresolved, including pricing guidance, site work, permits, architecture, land, design-build responsibility, and what marketing support actually includes.",
        "The plan separates demand generation from sales follow-up. A raised hand is not automatically a qualified opportunity, and the lead definition remains a leadership decision instead of a fabricated metric.",
    ):
        add_bullet(doc, text)

    add_callout(
        doc,
        "Evidence boundary",
        "Janice sent the recordings for analysis, but the formal public-use permission readback was not captured. Only the anonymized Homearama example has explicit paraphrase permission. Every other interview-derived public claim remains a factual-review draft.",
        "red",
    )

    # Commercial truth
    add_section_break(doc, "Commercial truth and scope", "A final proposal should make the approved commitment unmistakable and keep future work separate.")
    add_metric_strip(doc, [("35", "approved hours"), ("$30", "hourly rate"), ("$1,050", "approved total")])
    add_heading(doc, "Included in the current pilot", 2)
    for text in (
        "Audit of the current builder page and surrounding authority opportunity.",
        "Competitor and keyword research, content architecture, and AI-search strategy.",
        "One interview-integrated authority hub and two complete supporting articles.",
        "Implementation specification, metadata, internal-link logic, FAQ/schema candidate, and release QA gates.",
        "90-day roadmap, governed editorial workflow, final leadership packet, and bonus 12-month roadmap.",
    ):
        add_bullet(doc, text)

    add_heading(doc, "Mapped, but not produced in this pilot", 2)
    for text in (
        "Ongoing social post production.",
        "A broader downloadable asset beyond the mapped checklist and worksheet concepts.",
        "Video production or a webinar.",
        "The eleven next-phase articles, guides, subpillars, galleries, and tools in the authority roadmap.",
    ):
        add_bullet(doc, text)

    add_callout(doc, "Pricing boundary", "This packet does not invent a next-phase price. Cadence, production mix, ownership, revision allowance, payment timing, and acceptance mechanics must be scoped after leadership chooses the continuation model.", "gray")

    scope_rows = [
        ["Confirmed", "35 hours, $30/hour, $1,050 total, pilot deliverables, Phase 2 exclusions"],
        ["Decision required", "Invoice recipient and timing, payment timing, revision allowance, acceptance, ownership, confidentiality, portfolio use"],
        ["Not claimed", "Payment receipt, publication, traffic growth, ranking gains, lead volume, revenue, AI citations, or rich results"],
    ]
    add_table(doc, ["State", "Commercial position"], scope_rows, [1900, 7460], header_fill=INK, banded=True, font_size=9.2)

    # Content system
    add_section_break(doc, "The content authority system", "One commercial owner, distinct supporting jobs, and a complete path from discovery to qualified conversation.")
    add_heading(doc, "The hub-and-spoke model", 2)
    hub_num = new_numbering_instance(doc)
    add_numbered(doc, "Attract", "Capture builder-specific informational and commercial questions through distinct pages that match a real reader job.", hub_num)
    add_numbered(doc, "Prove", "Use original expertise, project photography, process evidence, client-approved examples, and direct answers that reduce uncertainty.", hub_num)
    add_numbered(doc, "Connect", "Link every supporting asset to the canonical builder hub and to the next logical decision page without creating duplicate intent.", hub_num)
    add_numbered(doc, "Convert", "Offer a low-pressure next step and preserve enough context for sales to evaluate fit, response time, and follow-up status.", hub_num)
    add_numbered(doc, "Learn", "Review actual queries, engagement, assisted actions, and lead-quality feedback before expanding the cluster.", hub_num)

    add_heading(doc, "Three production-ready pilot assets", 2)
    current_rows = []
    for asset in current_assets:
        current_rows.append([
            asset["asset_id"],
            asset["title"],
            f"{counts[asset['asset_id']]:,} words",
            asset["cta"],
        ])
    add_table(doc, ["ID", "Draft", "Approx. length", "Primary next step"], current_rows, [850, 4570, 1350, 2590], header_fill=ORANGE, banded=True, font_size=8.7)

    add_heading(doc, "Canonical ownership rules", 2)
    for rule in plan["cannibalization_rules"][:-1]:
        add_bullet(doc, rule)

    # Pilot content
    add_section_break(doc, "Pilot content: what is ready", "The copy is complete as a review draft; claims and publication remain gated.")
    pilot_sections = [
        ("HUB-01", "Custom Home Builder Marketing That Builds Trust Before the First Call", "Existing authority hub", "custom home builder marketing", "Schedule a non-sales call", [
            "Leads with excellent work and project photography instead of generic agency claims.",
            "Frames the marketing problem around trust, clarity, proof, discoverability, demand, and follow-up.",
            "Addresses referrals, an aging website, uneven pipeline, DIY marketing, and the 'work should speak for itself' objection.",
            "Preserves /marketing-agency-for-builders/ as the single commercial owner.",
        ]),
        ("ART-01", "Why Your Home Builder Website Is Not Generating Qualified Leads", "Supporting diagnostic article", "home builder website not generating leads", "Request a builder website clarity review", [
            "Separates low traffic, poor-fit traffic, weak proof, unclear scope, friction, and incomplete follow-up.",
            "Gives a corrective sequence instead of treating a redesign as the automatic answer.",
            "Supports the hub without competing for agency or service-page intent.",
            "Creates a natural bridge to the website-design subpillar and SEO checklist.",
        ]),
        ("ART-02", "5 Articles Every Custom Home Builder Blog Needs", "Supporting editorial article", "home builder blog", "Schedule a builder content planning call", [
            "Turns repeated buyer questions into five durable article types tied to actual decisions.",
            "Explains how one SME interview can support articles, FAQs, email, sales follow-up, social, and video planning.",
            "Keeps the broader content-marketing system for CNT-01 to prevent duplication.",
            "Ends with a planning conversation rather than a ranking promise.",
        ]),
    ]
    for idx, (asset_id, title, role, query, cta, points) in enumerate(pilot_sections):
        add_kicker(doc, f"{asset_id}  |  {role}")
        add_heading(doc, title, 2)
        add_body(doc, f"Primary query owner: {query}. Primary CTA: {cta}.", bold_lead="Primary query owner:")
        for point in points:
            add_bullet(doc, point)
        if idx < len(pilot_sections) - 1:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            p_pr = p._p.get_or_add_pPr()
            p_bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "8")
            bottom.set(qn("w:color"), "E1E1E1")
            p_bdr.append(bottom)
            p_pr.append(p_bdr)

    add_callout(doc, "Companion copy package", "The complete editable pillar, two complete articles, visible FAQ set, and matching FAQ schema candidate are included beside this proposal for line-by-line review and WordPress implementation.", "blue")

    # Keywords
    add_section_break(doc, "Keyword and intent strategy", "The 62-row ledger is a decision system, not a promise to publish 62 pages.")
    add_metric_strip(doc, [("62", "queries reviewed"), ("58", "assigned"), (str(reject_count), "rejected"), ("1", "canonical hub")])
    add_body(doc, "The keyword plan combines dated Semrush demand and difficulty evidence, dated Moz rank and URL evidence, live-site verification, funnel intent, reader job, and ownership. Blank evidence is treated as unknown, not as zero. The workbook preserves all raw fields and separates build decisions from source limitations.")

    add_heading(doc, "Priority clusters", 2)
    cluster_rows = [
        ["Core authority", "HUB-01", "custom home builder marketing; home builder marketing; agency/company/firm variations"],
        ["Visibility", "VIS-01 / LOCAL-01", "SEO for home builders; SEO for custom home builders; local SEO for home builders"],
        ["Website", "ART-01 / WEB-01 / GAL-01", "website not generating leads; website design; examples; best websites"],
        ["Strategy and audience", "STR-01 / AUD-01", "marketing plan; strategies; target market segments; new home builder marketing"],
        ["Lead operations", "LEAD-01 / NUR-01 / AUTO-01", "lead generation; CRM; email; nurture; automation; sales alignment"],
        ["Content", "ART-02 / CNT-01 / SOC-01", "home builder blog; content marketing; social media marketing"],
    ]
    add_table(doc, ["Cluster", "Owner", "Representative demand"], cluster_rows, [1800, 2000, 5560], header_fill=ORANGE, banded=True, font_size=8.7)

    add_heading(doc, "Rejected and validate-first demand", 2)
    add_bullet(doc, "Reject generic agency terms such as 'digital marketing agency' and 'SEO agency' for this focused pilot.")
    add_bullet(doc, "Reject homeowner-intent queries such as 'custom home builders near me' and 'how much does a custom home cost.'")
    add_bullet(doc, "Validate ambiguous phrases such as 'custom home builder website,' website classification criteria, and home builder market research before production.")
    add_bullet(doc, "Protect the existing hub's commercial relevance rather than splitting near-synonyms into thin competing pages.")

    # Portfolio roadmap
    add_section_break(doc, "The 19-asset portfolio", "Every keyword has an owner, every owner has a reader job, and every build window is explicit.")
    portfolio_rows = []
    for asset in assets:
        portfolio_rows.append([
            asset["asset_id"],
            asset["title"],
            asset["build_window"],
            asset["scope_state"],
        ])
    add_callout(doc, "Scope read", f"Three assets are complete current-pilot drafts. Eleven assets are mapped as next-phase production candidates. Social production is explicitly Phase 2, while the remaining assets stay in backlog or validate-first state.", "blue")
    add_table(doc, ["ID", "Asset", "Window", "State"], portfolio_rows, [820, 3900, 1770, 2870], header_fill=INK, banded=True, font_size=7.25)

    # 90 day
    add_section_break(doc, "90-day release and expansion plan", "The first 30 days are a governed release sequence. Expansion begins only after the foundation is verified.")
    release_rows = []
    for row in plan["calendar_90_day"]:
        release_rows.append([row["week"], row["asset"], row["scope"], row["status"]])
    add_table(doc, ["Timing", "Asset / gate", "Scope", "Current state"], release_rows, [1100, 3100, 1900, 3260], header_fill=ORANGE, banded=True, font_size=8.1)

    add_heading(doc, "Release logic", 2)
    release_num = new_numbering_instance(doc)
    for title, detail in (
        ("Gate 0", "Close Janice factual and permission review, leadership direction, imagery, owners, and the staging target."),
        ("Weeks 1-2", "Close copy, metadata, links, technical requirements, analytics, CRM, schema, accessibility, performance, and rollback preparation."),
        ("Weeks 3-5", "Assemble staging, run final QA, publish only the exact approved revision, then verify the live result separately."),
        ("Weeks 6-12", "Build the first next-phase subpillars only under a separately approved production scope and refreshed evidence."),
    ):
        add_numbered(doc, title, detail, release_num)

    add_callout(doc, "Expectation setting", "Ninety days is enough to launch a healthy foundation and learn from early demand signals. Six months is the more realistic window for momentum. This is a sequencing principle, not a traffic, lead, or revenue forecast.", "orange")

    # 12 month
    add_section_break(doc, "12-month authority roadmap", "A sustainable rhythm: one substantial asset, one meaningful refresh, one proof or tool, and one measurement question.")
    roadmap_rows = []
    for row in plan["roadmap_12_month"]:
        roadmap_rows.append([row["month"], row["focus"], row["primary_asset"], row["measurement"]])
    add_table(doc, ["Month", "Focus", "Primary asset", "Measurement question"], roadmap_rows, [750, 2420, 2300, 3890], header_fill=INK, banded=True, font_size=8.05)
    add_body(doc, "The roadmap is a planning model, not a production commitment. Each month requires an owner, approved scope, current source review, claim permissions, and release QA. Underperforming or overlapping assets should be refreshed, consolidated, or pruned rather than multiplied.")

    # SEO/AI implementation
    add_section_break(doc, "SEO, AI-search, and technical implementation", "Useful original content, clear ownership, and crawlable structure are the strategy; there is no separate AI-search shortcut.")
    add_heading(doc, "Implementation principles", 2)
    implementation_num = new_numbering_instance(doc)
    for title, detail in (
        ("Preserve one representative URL", "Keep the existing builder URL as the canonical commercial hub. Supporting assets own distinct reader jobs and link back to it."),
        ("Make expertise attributable", "Use named reviewers, first-hand observations, original visuals, dated sources, and explicit claim permissions."),
        ("Answer the whole decision", "Provide direct answers, scope clarity, process, proof, objections, and a logical next step instead of producing shallow query variants."),
        ("Keep machine-readable markup honest", "Use schema only when it matches visible content and the page type. Do not duplicate organization, breadcrumb, or FAQ markup."),
        ("Measure field performance", "Use Search Console and real-user Core Web Vitals data for actual evaluation. Lab tools guide diagnosis but are not proof of field experience."),
        ("Control crawlers intentionally", "Review search and training crawler policy as a business decision. Do not change robots.txt reflexively or imply that access guarantees inclusion."),
    ):
        add_numbered(doc, title, detail, implementation_num)

    add_heading(doc, "Page-level release specification", 2)
    for text in (
        "Confirm title, meta description, one H1, canonical, breadcrumbs, indexability, and sitemap behavior.",
        "Install reciprocal internal links and document the reader job and primary query owner.",
        "Match visible FAQ copy and FAQ schema exactly; do not promise a Google FAQ rich result.",
        "Verify desktop/mobile layout, keyboard and focus behavior, reduced motion, overflow, accessibility, forms, console output, analytics, and performance.",
        "Record a rollback snapshot and owner before publication, then verify the live rendered state after release.",
    ):
        add_bullet(doc, text)

    add_callout(doc, "No guarantee", "No ranking, traffic, lead, revenue, AI citation, answer-engine inclusion, featured snippet, or rich-result outcome is promised in this proposal.", "red")

    # Measurement
    add_section_break(doc, "Measurement and lead handoff", "The scorecard separates visibility, engagement, qualified actions, sales handling, and business outcomes.")
    measurement_rows = [
        ["Technical foundation", "Indexation, canonical selected by Google, sitemap status, crawl errors, Core Web Vitals", "Search Console, field data, release log", "After release and monthly"],
        ["Demand visibility", "Query mix, impressions, clicks, landing pages, nonbrand visibility", "Search Console and dated rank tools", "Monthly; no guaranteed target"],
        ["Content usefulness", "Engaged sessions, scroll depth where configured, CTA progression, internal-link paths", "Analytics and behavior instrumentation", "Monthly"],
        ["Qualified actions", "Consultation requests, calls, form completions, downloads, assisted actions", "Analytics and CRM reconciliation", "Weekly / monthly"],
        ["Sales handling", "Response time, acceptance, owner, status, nurture completion, disqualification reason", "CRM and sales-owner review", "Weekly"],
        ["Business outcome", "Qualified opportunities and assisted pipeline only after definitions and attribution are validated", "CRM and financial evidence", "Quarterly"],
    ]
    add_table(doc, ["Layer", "What to observe", "Source", "Cadence"], measurement_rows, [1550, 3520, 2600, 1690], header_fill=ORANGE, banded=True, font_size=8.05)

    add_heading(doc, "Qualified-lead definition to complete", 2)
    for text in (
        "Audience and market fit: builder type, geography, service area, and project type.",
        "Business fit: need, urgency, decision authority, current capacity, and willingness to invest.",
        "Engagement fit: raised-hand action, requested service, context supplied, and response status.",
        "Sales outcome: accepted, nurture, disqualified with reason, or no response after documented follow-up.",
    ):
        add_bullet(doc, text)
    add_callout(doc, "Open definition", "Janice described an MQL as a raised hand, but the full scoring model was not completed in the recorded interview. Leadership and sales must approve the definition before any conversion claim is used.", "orange")

    # Governance
    add_section_break(doc, "Editorial governance and permissions", "Interview insight becomes public content only through a documented review path.")
    add_heading(doc, "Claim lifecycle", 2)
    claim_num = new_numbering_instance(doc)
    for title, detail in (
        ("Capture", "Record the source, speaker, date, timestamp, working claim, and intended asset."),
        ("Classify", "Mark the claim as experience, opinion, process, terminology, estimate, or client example."),
        ("Verify", "Check every number, name, result, award, ranking, timeframe, and platform behavior against an original source."),
        ("Permission", "Assign an explicit public-use code and required anonymization. Silence is not permission."),
        ("Edit", "Convert approved evidence into precise reader-facing language without adding magnitude or certainty."),
        ("Release", "Publish only the exact reviewed revision, then verify the live page separately."),
    ):
        add_numbered(doc, title, detail, claim_num)

    add_heading(doc, "Interview material currently held", 2)
    held_items = [
        "Audience boundaries across custom builders, production builders, developers, master-planned communities, and remodelers.",
        "Any revenue or business-maturity threshold.",
        "Trigger moments, trust factors, photography guidance, website questions, referral limits, objections, and terminology.",
        "The full qualified-lead definition and measurement language.",
    ]
    for item in held_items:
        add_bullet(doc, item)
    add_callout(doc, "Explicit permission", "The anonymized Homearama example may be paraphrased. The visitor estimate, named client, timeframe, performance implication, and any other result still require independent verification before use.", "green")

    # Decisions
    add_section_break(doc, "September 3 leadership decisions", "The system is ready for a bounded approval conversation, not an open-ended rewrite.")
    decision_rows = [
        ["1", "Visual direction", "Choose Cinematic Authority, Orange Press, or an explicitly bounded hybrid."],
        ["2", "Canonical owner", "Confirm /marketing-agency-for-builders/ remains the one authority hub."],
        ["3", "Approvers", "Name owners for copy, claims, visuals, technical changes, publication, and rollback."],
        ["4", "Staging", "Confirm the WordPress staging target, production target, access path, and rollback owner."],
        ["5", "Imagery", "Approve hero, project photography, article featured images, releases, and alt-text owner."],
        ["6", "Measurement", "Approve the scorecard and complete the qualified-lead definition."],
        ["7", "Continuation", "Choose next-phase cadence, deliverable mix, ownership, and production scope."],
        ["8", "Commercial terms", "Confirm invoice/payment timing, revisions, acceptance, ownership, confidentiality, and portfolio use."],
    ]
    add_table(doc, ["#", "Decision", "Required resolution"], decision_rows, [520, 2050, 6790], header_fill=INK, banded=True, font_size=8.65)

    add_heading(doc, "Janice review checklist", 2)
    for text in (
        "Confirm audience boundaries and remove or confirm the tentative maturity threshold.",
        "Approve the edited positioning, trust, photography, website, referral, pricing-guidance, and objection language.",
        "Approve five to eight visible FAQ answers and matching schema candidate.",
        "Complete the qualified-lead definition.",
        "Approve or reject the anonymized Homearama paraphrase and assign public-use codes to all other interview language.",
    ):
        add_bullet(doc, text)

    # Acceptance
    add_section_break(doc, "Acceptance and release gates", "Completion of the proposal is not the same as approval, publication, or verified live performance.")
    add_heading(doc, "Pilot package acceptance", 2)
    for text in (
        "Leadership confirms the packet reflects the approved 35-hour pilot and no additional production has been implied.",
        "Janice confirms factual accuracy and permission states for interview-derived language.",
        "Leadership selects the visual direction and exact WordPress revision path.",
        "Commercial mechanics are documented: invoice, payment, revisions, acceptance, ownership, confidentiality, and portfolio use.",
        "The content workbook and companion drafts are accepted as the operating source package for implementation.",
    ):
        add_bullet(doc, text)

    add_heading(doc, "Publication gate", 2)
    for text in (
        "Approved copy and imagery are installed in staging.",
        "Metadata, canonical, breadcrumbs, internal links, sitemap behavior, and schema are verified without duplication.",
        "Visible FAQ copy and FAQ schema match exactly.",
        "Desktop/mobile, keyboard/focus, reduced motion, overflow, accessibility, console, links, forms, analytics, and performance pass final QA.",
        "Rollback snapshot and owner are recorded.",
        "Publication is approved for the exact revision, then the live page is reverified independently.",
    ):
        add_bullet(doc, text)

    add_callout(doc, "Current status", "Proposal, workbook, content drafts, and operating plan are complete. Janice review, leadership decisions, WordPress staging, publication approval, and live verification remain open.", "blue")

    # Sources
    add_section_break(doc, "Evidence and source notes", "Current verification, dated exports, professional opinion, and approval-gated interview evidence are kept separate.")
    source_rows = [
        ["Janice call recordings", "Aug. 28, 2026", "Audience, objections, trust, website, timing, example", "Most public permission unresolved"],
        ["Approved paid-pilot proposal", "Reconciled Aug. 3, 2026", "Scope, 35 hours, $30/hour, $1,050", "Commercial mechanics open"],
        ["Semrush keyword map", "Aug. 4, 2026", "Demand, intent, difficulty, CPC", "Dated; not a forecast"],
        ["Moz rankings export", "Through Aug. 14, 2026", "Protect existing URL ownership", "Not a current live rank check"],
        ["Current builder page", "Verified Aug. 28, 2026", "HTTP 200, self-canonical, live copy, JSON-LD", "Current state; not replacement approval"],
        ["Google Search documentation", "Verified Aug. 28, 2026", "AI features, people-first content, canonical, structured data, FAQ", "Official guidance; no outcome promise"],
        ["Web Vitals guidance", "Verified Aug. 28, 2026", "Field-performance measurement", "Use real-user data for pass/fail"],
    ]
    add_table(doc, ["Source", "Freshness", "Use", "Boundary"], source_rows, [1900, 1600, 3150, 2710], header_fill=ORANGE, banded=True, font_size=8.0)

    add_heading(doc, "Official implementation references", 2)
    refs = [
        "Google Search: Optimizing your website for generative AI features on Google Search",
        "Google Search: Creating helpful, reliable, people-first content",
        "Google Search: Guidance about AI-generated content",
        "Google Search: Canonicalization and duplicate URL guidance",
        "Google Search: Structured data policies and FAQ rich-result changes",
        "Google Search: SEO Starter Guide",
        "Web.dev: Web Vitals",
        "OpenAI and Perplexity crawler-control guidance",
    ]
    for ref in refs:
        add_bullet(doc, ref)

    add_callout(doc, "Final position", "Approve the governed foundation first. Publish one coherent hub and two distinct supporting articles. Learn from actual queries and qualified actions. Expand only when the evidence, ownership, and production scope justify the next asset.", "orange")

    # Final metadata / save
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(json.dumps({
        "output": str(OUT),
        "assets": len(assets),
        "keywords": len(plan["keyword_owner"]),
        "pilot_word_counts": counts,
    }, indent=2))


if __name__ == "__main__":
    build()
