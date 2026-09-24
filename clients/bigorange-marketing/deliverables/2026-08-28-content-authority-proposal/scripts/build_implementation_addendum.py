from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

ROOT = Path(__file__).resolve().parent.parent
RELEASE = ROOT / "release"
OUT = RELEASE / "leadership-addendum"
OUT.mkdir(parents=True, exist_ok=True)
PLAN = json.loads((ROOT / "working" / "content-plan.json").read_text(encoding="utf-8"))
QA_PATH = RELEASE / "qa" / "content-system-qa.json"
QA = json.loads(QA_PATH.read_text(encoding="utf-8")) if QA_PATH.exists() else {"status": "not-run", "asset_count": 0, "error_count": None}

DOCX = OUT / "BigOrange-19-Asset-Content-Production-Addendum-2026-08-28.docx"
PDF = OUT / "BigOrange-19-Asset-Content-Production-Addendum-2026-08-28.pdf"

INK = RGBColor(22, 30, 42)
ORANGE = RGBColor(246, 112, 42)
MUTED = RGBColor(92, 103, 117)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def docx_build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    for name, size, color in [("Title", 25, INK), ("Heading 1", 16, INK), ("Heading 2", 12, ORANGE)]:
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = color
        styles[name].font.bold = True

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("BIGORANGE AUTHORITY SYSTEM")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(10)
    run.font.color.rgb = ORANGE
    title = doc.add_paragraph("19-Asset Content Production Addendum", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub = doc.add_paragraph("Full SEO, AEO, and GEO review package | August 28, 2026")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].font.color.rgb = MUTED

    doc.add_heading("Bottom line", level=1)
    doc.add_paragraph(
        "All 19 owned URL assets now have complete local review drafts. The release also includes WordPress-ready HTML fragments, one candidate schema graph per URL, reciprocal internal-link specifications, metadata, image direction, source notes, and machine-readable QA. Nothing has been sent, staged, published, or represented as live performance."
    )
    facts = doc.add_table(rows=2, cols=5)
    facts.style = "Table Grid"
    labels = ["Keyword decisions", "Assigned", "Rejected", "Owned assets", "QA state"]
    vals = ["62", "58", "4", "19", str(QA.get("status", "not-run")).upper()]
    for i, value in enumerate(labels):
        facts.cell(0, i).text = value
        set_cell_shading(facts.cell(0, i), "161E2A")
        for run in facts.cell(0, i).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
        facts.cell(1, i).text = vals[i]

    doc.add_heading("Scope truth", level=1)
    for text in [
        "The original approved pilot remains 35 hours at $30 per hour, $1,050 total.",
        "The original pilot included three complete URL drafts: HUB-01, ART-01, and ART-02.",
        "The remaining 16 complete local review drafts are a Dillon-authorized production extension. They do not change the original commercial record or imply a new client fee.",
        "Janice factual and public-use review, leadership decisions, imagery, staging ownership, exact-revision approval, and live verification remain gates.",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Complete asset register", level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["ID", "Asset", "Reader job", "Original window", "Current local state"]
    for i, header in enumerate(headers):
        table.cell(0, i).text = header
        set_cell_shading(table.cell(0, i), "161E2A")
        for run in table.cell(0, i).paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
    for asset in PLAN["assets"]:
        row = table.add_row().cells
        values = [asset["asset_id"], asset["title"], asset["reader_job"], asset["build_window"], "Complete local review draft"]
        for i, value in enumerate(values):
            row[i].text = value

    doc.add_page_break()
    doc.add_heading("What is implementation-ready", level=1)
    ready = [
        "Complete source Markdown for all 19 URL owners",
        "Title tags, meta descriptions, slugs, canonical candidates, primary and secondary topics",
        "Direct-answer introductions, extractable tables or steps, authoritative source links, and internal links",
        "Visible FAQs and matching FAQ schema candidates for the 16 new drafts; approved HUB-01 companion candidate preserved",
        "Featured-image briefs, alt-text candidates, and inline image placement notes",
        "WordPress-ready HTML fragments, schema files, publication manifest, and reciprocal link map",
        "Automated QA output and a WordPress staging and live-release checklist",
    ]
    for item in ready:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("What remains human-gated", level=1)
    gated = [
        "Janice approval of factual accuracy and public-use permission for interview-derived language",
        "Leadership approval of audience, positioning, offers, claims, imagery, CTA destinations, approvers, and qualified-lead definition",
        "A distinct consulting offer before CONSULT-01 can become a public service page",
        "The exact WordPress staging target, editor, analytics owner, CRM owner, publication owner, and rollback owner",
        "Final staged accessibility, mobile, link, form, metadata, schema, analytics, CRM, performance, and rollback QA",
        "Explicit publication approval for the exact staged revision and separate live readback",
    ]
    for item in gated:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("No outcome claim", level=1)
    doc.add_paragraph(
        "Completion of this local production package does not guarantee or claim rankings, traffic, leads, revenue, featured snippets, rich results, AI citations, answer-engine inclusion, publication, or live performance."
    )
    doc.save(DOCX)


def pdf_build() -> None:
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="BO_Title", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=24, leading=28, textColor=colors.HexColor("#161E2A"), alignment=TA_CENTER, spaceAfter=12))
    styles.add(ParagraphStyle(name="BO_Sub", parent=styles["Normal"], fontName="Helvetica", fontSize=10, textColor=colors.HexColor("#5C6775"), alignment=TA_CENTER, spaceAfter=20))
    styles.add(ParagraphStyle(name="BO_H1", parent=styles["Heading1"], fontName="Helvetica-Bold", fontSize=16, leading=19, textColor=colors.HexColor("#161E2A"), spaceBefore=10, spaceAfter=8))
    styles.add(ParagraphStyle(name="BO_Body", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=13, textColor=colors.HexColor("#253142"), spaceAfter=6))
    styles.add(ParagraphStyle(name="BO_Small", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.2, leading=9, textColor=colors.HexColor("#253142")))
    doc = SimpleDocTemplate(str(PDF), pagesize=landscape(letter), rightMargin=0.45 * inch, leftMargin=0.45 * inch, topMargin=0.45 * inch, bottomMargin=0.45 * inch)
    story = [
        Paragraph("BIGORANGE AUTHORITY SYSTEM", ParagraphStyle(name="eyebrow", parent=styles["BO_Sub"], textColor=colors.HexColor("#F6702A"), fontName="Helvetica-Bold", spaceAfter=5)),
        Paragraph("19-Asset Content Production Addendum", styles["BO_Title"]),
        Paragraph("Full SEO, AEO, and GEO review package | August 28, 2026", styles["BO_Sub"]),
        Paragraph("Bottom line", styles["BO_H1"]),
        Paragraph("All 19 owned URL assets now have complete local review drafts. The release includes WordPress-ready HTML fragments, candidate schema graphs, reciprocal internal-link specifications, metadata, image direction, source notes, and automated QA. Nothing has been sent, staged, published, or represented as live performance.", styles["BO_Body"]),
    ]
    facts = [["Keyword decisions", "Assigned", "Rejected", "Owned assets", "QA state"], ["62", "58", "4", "19", str(QA.get("status", "not-run")).upper()]]
    fact_table = Table(facts, colWidths=[1.55 * inch] * 5)
    fact_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#161E2A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D7DDE4")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([fact_table, Spacer(1, 12), Paragraph("Scope truth", styles["BO_H1"])])
    for text in [
        "The original approved pilot remains 35 hours at $30 per hour, $1,050 total.",
        "The original pilot included HUB-01, ART-01, and ART-02 as complete URL drafts.",
        "The other 16 complete drafts are a Dillon-authorized local production extension. They do not change the original commercial record or imply a new client fee.",
        "Factual review, public-use permission, leadership approval, imagery, staging, exact-revision approval, and live verification remain gates.",
    ]:
        story.append(Paragraph(f"• {text}", styles["BO_Body"]))
    story.extend([PageBreak(), Paragraph("Complete 19-asset register", styles["BO_H1"])])
    rows = [["ID", "Asset", "Original window", "Current local state"]]
    for asset in PLAN["assets"]:
        rows.append([asset["asset_id"], Paragraph(asset["title"], styles["BO_Small"]), Paragraph(asset["build_window"], styles["BO_Small"]), "Complete review draft"])
    asset_table = Table(rows, colWidths=[0.75 * inch, 4.2 * inch, 2.4 * inch, 2.1 * inch], repeatRows=1)
    asset_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#161E2A")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTNAME", (0, 1), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.2),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#D7DDE4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F8FA")]),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(asset_table)
    story.extend([PageBreak(), Paragraph("Implementation-ready package", styles["BO_H1"])])
    for item in [
        "Complete Markdown copy for all 19 owners, with metadata and one distinct reader job per URL.",
        "Direct-answer openings, extractable structures, authoritative sources, planned internal links, FAQs, and schema candidates.",
        "Featured-image briefs, alt-text candidates, and image placement notes.",
        "WordPress-ready HTML, per-URL schema files, publication manifest, internal-link map, and automated QA.",
        "AEO, GEO, SEO, WordPress staging, publication, and live-verification standards.",
    ]:
        story.append(Paragraph(f"• {item}", styles["BO_Body"]))
    story.append(Paragraph("Human gates still open", styles["BO_H1"]))
    for item in [
        "Janice factual and public-use review for interview-derived language.",
        "Leadership approval of audience, claims, offers, visuals, CTA destinations, owners, and qualified-lead definition.",
        "A distinct consulting offer before CONSULT-01 can become a public service page.",
        "Exact WordPress staging and rollback owners, followed by complete staged QA.",
        "Explicit publication approval for the exact staged revision and separate live readback.",
    ]:
        story.append(Paragraph(f"• {item}", styles["BO_Body"]))
    story.append(Paragraph("No outcome claim", styles["BO_H1"]))
    story.append(Paragraph("This local production package does not guarantee or claim rankings, traffic, leads, revenue, featured snippets, rich results, AI citations, answer-engine inclusion, publication, or live performance.", styles["BO_Body"]))
    doc.build(story)


docx_build()
pdf_build()
print(json.dumps({"docx": str(DOCX), "pdf": str(PDF), "assets": len(PLAN["assets"]), "qa": QA.get("status")}, indent=2))
