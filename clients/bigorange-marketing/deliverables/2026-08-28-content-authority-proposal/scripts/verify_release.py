from __future__ import annotations

import csv
import hashlib
import html
import json
import re
import zipfile
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent.parent
RELEASE = ROOT / "release"
QA = RELEASE / "qa"
ZIP = RELEASE / "BigOrange-19-Asset-AEO-GEO-SEO-Content-Package-2026-08-28.zip"
WORKBOOK = ROOT / "outputs" / "01a04906-96f3-7072-b3bf-0c2589aed6fa" / "BigOrange-19-Asset-Content-Production-Plan-2026-08-28.xlsx"
PDF = RELEASE / "leadership-addendum" / "BigOrange-19-Asset-Content-Production-Addendum-2026-08-28.pdf"
DOCX = RELEASE / "leadership-addendum" / "BigOrange-19-Asset-Content-Production-Addendum-2026-08-28.docx"


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_html_text(value: str) -> str:
    value = re.sub(r"<[^>]+>", " ", value)
    return re.sub(r"\s+", " ", html.unescape(value)).strip()


def visible_html_faq_pairs(document: str) -> list[tuple[str, str]]:
    section = re.search(
        r"<h2[^>]*>\s*Frequently asked questions\s*</h2>(.*?)(?=<h2\b|\Z)",
        document,
        re.I | re.S,
    )
    if not section:
        return []
    return [
        (clean_html_text(question), clean_html_text(answer))
        for question, answer in re.findall(
            r"<h3[^>]*>(.*?)</h3>\s*<p[^>]*>(.*?)</p>",
            section.group(1),
            re.I | re.S,
        )
    ]


def schema_faq_pairs(data: dict) -> list[tuple[str, str]]:
    for item in data.get("@graph", []):
        if isinstance(item, dict) and item.get("@type") == "FAQPage":
            pairs = []
            for entity in item.get("mainEntity", []):
                answer = entity.get("acceptedAnswer", {})
                pairs.append((
                    re.sub(r"\s+", " ", str(entity.get("name", ""))).strip(),
                    re.sub(r"\s+", " ", str(answer.get("text", ""))).strip(),
                ))
            return pairs
    return []


errors: list[str] = []
warnings: list[str] = []
evidence: dict = {}

source_files = sorted((RELEASE / "source-markdown").glob("*.md"))
html_files = sorted((RELEASE / "wordpress-ready-html").glob("*.html"))
schema_files = sorted((RELEASE / "schema").glob("*.json"))
evidence["source_markdown_files"] = len(source_files)
evidence["html_fragments"] = len(html_files)
evidence["schema_files"] = len(schema_files)
if len(source_files) != 20:
    errors.append(f"Expected 20 Markdown files including the HUB FAQ companion, found {len(source_files)}")
if len(html_files) != 19:
    errors.append(f"Expected 19 HTML fragments, found {len(html_files)}")
if len(schema_files) != 19:
    errors.append(f"Expected 19 schema files, found {len(schema_files)}")

schema_summary = []
for path in schema_files:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except Exception as exc:
        errors.append(f"Malformed schema {path.name}: {exc}")
        continue
    graph = data.get("@graph", [])
    types = [item.get("@type") for item in graph if isinstance(item, dict)]
    if "Article" not in types:
        errors.append(f"{path.name} lacks Article in @graph")
    asset_id = path.name.split(".", 1)[0]
    matching_html = next((item for item in html_files if item.name.startswith(f"{asset_id}-")), None)
    schema_pairs = schema_faq_pairs(data)
    if schema_pairs:
        if not matching_html:
            errors.append(f"{path.name} has FAQPage but no matching HTML fragment")
        else:
            visible_pairs = visible_html_faq_pairs(matching_html.read_text(encoding="utf-8"))
            if visible_pairs != schema_pairs:
                errors.append(f"{asset_id}: visible HTML FAQ does not exactly match FAQPage schema")
    schema_summary.append({"file": path.name, "types": types})
evidence["schema_summary"] = schema_summary

with (RELEASE / "publication-manifest.csv").open(encoding="utf-8-sig", newline="") as stream:
    manifest = list(csv.DictReader(stream))
with (RELEASE / "internal-link-map.csv").open(encoding="utf-8-sig", newline="") as stream:
    links = list(csv.DictReader(stream))
evidence["publication_manifest_rows"] = len(manifest)
evidence["internal_link_rows"] = len(links)
if len(manifest) != 19:
    errors.append(f"Publication manifest has {len(manifest)} rows, expected 19")
if not links:
    errors.append("Internal link map is empty")
manifest_by_id = {row["asset_id"]: row for row in manifest}
for row_number, row in enumerate(links, start=2):
    source_id = row["source_asset_id"]
    slug = row["target_slug"]
    source = manifest_by_id.get(source_id)
    if not source:
        errors.append(f"Internal-link row {row_number} references unknown source {source_id}")
        continue
    source_text = (ROOT / source["source_file"]).read_text(encoding="utf-8")
    html_text = (RELEASE / source["html_fragment"]).read_text(encoding="utf-8")
    if f"]({slug})" not in source_text:
        errors.append(f"Internal-link row {row_number}: {source_id} source lacks {slug}")
    if f'href="{slug}"' not in html_text:
        errors.append(f"Internal-link row {row_number}: {source_id} HTML lacks {slug}")

content_qa = json.loads((QA / "content-system-qa.json").read_text(encoding="utf-8"))
evidence["content_qa"] = {k: content_qa.get(k) for k in ["status", "asset_count", "error_count", "warning_count"]}
if content_qa.get("status") != "pass" or content_qa.get("asset_count") != 19 or content_qa.get("error_count") != 0:
    errors.append("Content-system QA is not a clean 19-asset pass")

wb = load_workbook(WORKBOOK, read_only=False, data_only=False)
evidence["workbook_sheets"] = wb.sheetnames
if len(wb.sheetnames) != 8:
    errors.append(f"Workbook has {len(wb.sheetnames)} sheets, expected 8")
formula_error_tokens = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")
formula_errors = []
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and any(token in cell.value for token in formula_error_tokens):
                formula_errors.append(f"{ws.title}!{cell.coordinate}={cell.value}")
if formula_errors:
    errors.append(f"Workbook formula error tokens found: {formula_errors[:10]}")
wb.close()

pdf_reader = PdfReader(str(PDF))
doc = Document(DOCX)
evidence["addendum_pdf_pages"] = len(pdf_reader.pages)
evidence["addendum_docx_paragraphs"] = len(doc.paragraphs)
if len(pdf_reader.pages) != 3:
    errors.append(f"Leadership addendum PDF has {len(pdf_reader.pages)} pages, expected 3")
if len(doc.paragraphs) < 10:
    errors.append("Leadership addendum DOCX did not reopen with expected content")

if not ZIP.exists():
    errors.append("Release ZIP is missing")
else:
    with zipfile.ZipFile(ZIP) as archive:
        bad = archive.testzip()
        names = archive.namelist()
    evidence["zip_entries"] = len(names)
    evidence["zip_bad_entry"] = bad
    if bad:
        errors.append(f"ZIP CRC failure: {bad}")
    if sum(1 for n in names if n.lower().endswith(".html")) != 19:
        errors.append("ZIP does not contain 19 HTML fragments")
    if sum(1 for n in names if "/schema/" in f"/{n.lower()}" and n.lower().endswith(".json")) != 19:
        errors.append("ZIP does not contain 19 schema JSON files")
    if any(n.lower().endswith((".m4a", ".mp3", ".wav", ".part2")) for n in names):
        errors.append("ZIP unexpectedly contains a source recording")
    if any(n.lower().endswith(("qa/release-integrity.json", "qa/release-integrity.md")) for n in names):
        errors.append("ZIP embeds the self-referential release-integrity sidecar")

critical = [
    RELEASE / "CONTENT-SYSTEM-INDEX.md",
    RELEASE / "publication-manifest.csv",
    RELEASE / "internal-link-map.csv",
    WORKBOOK,
    PDF,
    DOCX,
    ZIP,
]
evidence["hashes"] = {str(path.relative_to(ROOT)): sha256(path) for path in critical if path.exists()}

result = {
    "checked_at": "2026-08-28",
    "status": "pass" if not errors else "fail",
    "errors": errors,
    "warnings": warnings,
    "evidence": evidence,
}
QA.mkdir(parents=True, exist_ok=True)
(QA / "release-integrity.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
summary = [
    "# Release Integrity",
    "",
    f"Status: **{result['status'].upper()}**",
    f"Errors: {len(errors)}",
    f"Warnings: {len(warnings)}",
    "",
    f"- Editable Markdown files: {len(source_files)} (19 URL owners plus HUB FAQ companion)",
    f"- WordPress HTML fragments: {len(html_files)}",
    f"- Schema graphs: {len(schema_files)}",
    f"- Publication rows: {len(manifest)}",
    f"- Internal-link rows: {len(links)}",
    f"- Workbook sheets: {len(evidence.get('workbook_sheets', []))}",
    f"- Addendum PDF pages: {evidence.get('addendum_pdf_pages')}",
    f"- ZIP entries: {evidence.get('zip_entries')}",
    "- Release-integrity JSON/Markdown: external sidecars, intentionally excluded from the ZIP they validate",
]
if errors:
    summary += ["", "## Errors", ""] + [f"- {item}" for item in errors]
(QA / "release-integrity.md").write_text("\n".join(summary) + "\n", encoding="utf-8")
print(json.dumps({"status": result["status"], "errors": len(errors), "warnings": len(warnings), "zip_entries": evidence.get("zip_entries")}, indent=2))
raise SystemExit(0 if not errors else 1)
