from __future__ import annotations

import math
import os
import re
import zipfile
from pathlib import Path

from fontTools.ttLib import TTFont as FontToolsTTFont
from openpyxl import load_workbook
from reportlab.lib.colors import HexColor, Color
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader


ROOT = Path(__file__).resolve().parent
SOURCE = next(
    path
    for path in (
        ROOT / "upload" / "Align_HCM_Top30_Story_Shortlist_Branded.xlsx",
        ROOT / "Align_HCM_Top30_Story_Shortlist_Branded.xlsx",
    )
    if path.exists()
)
TMP = ROOT / "tmp" / "pdfs"
OUTPUT = ROOT / "output" / "pdf" / "Align_HCM_Top30_Story_Shortlist_Editorial.pdf"
CLIENT_LOGO_DIR = TMP / "client-logos"

PAGE_W, PAGE_H = landscape(A4)
MARGIN = 34

NAVY = HexColor("#0A1628")
ORANGE = HexColor("#F05A28")
BRIGHT_ORANGE = HexColor("#FF6B2B")
RUST = HexColor("#AD3D1B")
TEAL = HexColor("#136E61")
PAPER = HexColor("#FCFAF7")
PANEL = HexColor("#F5F1EA")
HAIRLINE = HexColor("#E9E4DC")
SLATE = HexColor("#2D3748")
MUTED = HexColor("#646E7C")
BLUE = HexColor("#1B4F72")
BLUE_TINT = HexColor("#EDF3F8")
TEAL_TINT = HexColor("#E9F7F4")
RUST_TINT = HexColor("#FDF1EA")
YELLOW = HexColor("#FFF2CC")
WHITE = HexColor("#FFFFFF")


def ascii_text(value) -> str:
    if value is None:
        return ""
    text = str(value)
    return (
        text.replace("\u2014", " - ")
        .replace("\u2013", "-")
        .replace("\u2011", "-")
        .replace("\u2010", "-")
        .replace("\u00a0", " ")
        .replace("\u00b7", "|")
        .replace("\u2248", "about ")
    )


def convert_font(src: Path, dst: Path) -> None:
    if dst.exists():
        return
    font = FontToolsTTFont(str(src))
    font.flavor = None
    font.save(str(dst))


def register_fonts() -> None:
    TMP.mkdir(parents=True, exist_ok=True)
    font_root = Path("/tmp/align-top30-runtime/node_modules/@fontsource")
    candidates = {
        "Inter": font_root / "inter/files/inter-latin-400-normal.woff",
        "Inter-Medium": font_root / "inter/files/inter-latin-500-normal.woff",
        "Inter-SemiBold": font_root / "inter/files/inter-latin-600-normal.woff",
        "Inter-Bold": font_root / "inter/files/inter-latin-700-normal.woff",
        "Syne-Bold": font_root / "syne/files/syne-latin-700-normal.woff",
    }
    fallback = {
        "Inter": Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        "Inter-Medium": Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        "Inter-SemiBold": Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        "Inter-Bold": Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        "Syne-Bold": Path("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    }
    for name, src in candidates.items():
        try:
            if src.exists():
                ttf = TMP / f"{name}.ttf"
                convert_font(src, ttf)
                pdfmetrics.registerFont(TTFont(name, str(ttf)))
            else:
                pdfmetrics.registerFont(TTFont(name, str(fallback[name])))
        except Exception:
            pdfmetrics.registerFont(TTFont(name, str(fallback[name])))


def extract_align_logo() -> Path:
    out = TMP / "align-hcm-logo-reverse.png"
    if not out.exists():
        with zipfile.ZipFile(SOURCE) as zf:
            out.write_bytes(zf.read("xl/media/image1.png"))
    return out


def load_data():
    wb_values = load_workbook(SOURCE, data_only=True)
    wb_formulas = load_workbook(SOURCE, data_only=False)

    top = wb_values["Top 30 Shortlist"]
    headers = [ascii_text(top.cell(7, col).value) for col in range(1, 21)]
    stories = []
    for row in range(8, 38):
        values = [top.cell(row, col).value for col in range(1, 21)]
        stories.append(dict(zip(headers, values)))

    totals = {
        "value": float(top["J38"].value or sum(float(s["Value (true USD)"] or 0) for s in stories)),
        "narratives": int(top["K38"].value or sum(s["Narrative?"] == "Yes" for s in stories)),
        "quotes": int(top["L38"].value or sum(s["Quote?"] == "Raven" for s in stories)),
        "industries": len({s["Industry"] for s in stories}),
    }

    ref = wb_values["Reference Candidates"]
    ref_headers = [ascii_text(ref.cell(5, col).value) for col in range(1, 11)]
    references = []
    for row in range(6, ref.max_row + 1):
        values = [ref.cell(row, col).value for col in range(1, 11)]
        references.append(dict(zip(ref_headers, values)))

    gaps_ws = wb_values["Representation Gaps"]
    gap_sections = []
    current = None
    for row in range(5, gaps_ws.max_row + 1):
        vals = [gaps_ws.cell(row, col).value for col in range(1, 6)]
        if vals[0] and all(v is None for v in vals[1:]):
            current = {"title": ascii_text(vals[0]), "headers": [], "rows": []}
            gap_sections.append(current)
        elif current and vals[0] == "Value":
            current["headers"] = [ascii_text(v) for v in vals]
        elif current and vals[0] is not None:
            current["rows"].append(vals)

    readme = wb_formulas["Read Me"]
    method_sections = []
    current = None
    headings = {
        "Source & scope",
        "One correction you should apply to the source workbook",
        "Selection method",
        "What the shortlist needs from you",
        "Unblocking the other 421 engagements",
        "Branded edition, 10 Aug 2026",
    }
    for row in range(6, readme.max_row + 1):
        value = readme.cell(row, 1).value
        if not value:
            continue
        text = ascii_text(value)
        if value in headings:
            current = {"title": text, "paragraphs": []}
            method_sections.append(current)
        elif current:
            current["paragraphs"].append(text)

    return stories, totals, references, gap_sections, method_sections


def set_font(c: canvas.Canvas, name="Inter", size=10, color=SLATE):
    c.setFont(name, size)
    c.setFillColor(color)


def round_rect(c, x, y, w, h, fill, radius=12, stroke=None, line_width=0.8):
    c.saveState()
    c.setFillColor(fill)
    if stroke:
        c.setStrokeColor(stroke)
        c.setLineWidth(line_width)
    else:
        c.setStrokeColor(fill)
    c.roundRect(x, y, w, h, radius, fill=1, stroke=1 if stroke else 0)
    c.restoreState()


def wrap_lines(text, font, size, width, max_lines=None):
    text = ascii_text(text)
    if not text:
        return []
    words = text.split()
    lines, line = [], ""
    for word in words:
        candidate = word if not line else f"{line} {word}"
        if pdfmetrics.stringWidth(candidate, font, size) <= width:
            line = candidate
        else:
            if line:
                lines.append(line)
            line = word
            if max_lines and len(lines) >= max_lines:
                break
    if line and (not max_lines or len(lines) < max_lines):
        lines.append(line)
    if max_lines and len(lines) == max_lines and " ".join(lines) != text:
        while lines[-1] and pdfmetrics.stringWidth(lines[-1] + "...", font, size) > width:
            lines[-1] = lines[-1][:-1]
        lines[-1] = lines[-1].rstrip() + "..."
    return lines


def draw_wrapped(c, text, x, y, width, font="Inter", size=9, leading=None, color=SLATE, max_lines=None):
    leading = leading or size * 1.25
    set_font(c, font, size, color)
    lines = wrap_lines(text, font, size, width, max_lines)
    for line in lines:
        c.drawString(x, y, line)
        y -= leading
    return y


def draw_logo(c, logo_path, x, y, width):
    img = ImageReader(str(logo_path))
    iw, ih = img.getSize()
    height = width * ih / iw
    c.drawImage(img, x, y, width=width, height=height, mask="auto", preserveAspectRatio=True)
    return height


def logo_filename(story):
    domain = ascii_text(story.get("Website")) or ascii_text(story.get("Client"))
    safe = re.sub(r"[^a-z0-9]+", "-", domain.lower()).strip("-")
    return CLIENT_LOGO_DIR / f"{safe}.png"


def draw_client_logo(c, story, x, y, w, h, show_name_fallback=True):
    round_rect(c, x, y, w, h, WHITE, 8, HAIRLINE)
    path = logo_filename(story)
    if path.exists():
        try:
            img = ImageReader(str(path))
            iw, ih = img.getSize()
            scale = min((w - 12) / iw, (h - 10) / ih)
            dw, dh = iw * scale, ih * scale
            c.drawImage(
                img,
                x + (w - dw) / 2,
                y + (h - dh) / 2,
                width=dw,
                height=dh,
                mask="auto",
                preserveAspectRatio=True,
            )
            return True
        except Exception:
            pass
    if show_name_fallback:
        initials = "".join(word[0] for word in ascii_text(story.get("Client")).split()[:3]).upper()
        set_font(c, "Syne-Bold", 10, NAVY)
        c.drawCentredString(x + w / 2, y + h / 2 - 3.5, initials[:3])
    return False


def draw_header(c, title, section, logo_path, dark=False):
    bg = NAVY if dark else PAPER
    c.setFillColor(bg)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
    if dark:
        c.setFillColor(Color(1, 1, 1, alpha=0.06))
        c.circle(PAGE_W - 36, PAGE_H - 28, 144, fill=1, stroke=0)
        c.setFillColor(Color(1, 0.42, 0.17, alpha=0.14))
        c.circle(PAGE_W - 28, 46, 118, fill=1, stroke=0)
        draw_logo(c, logo_path, PAGE_W - 172, PAGE_H - 76, 132)
        set_font(c, "Inter-SemiBold", 9, HexColor("#B9C6D8"))
        c.drawString(MARGIN, PAGE_H - 38, section.upper())
        set_font(c, "Syne-Bold", 27, WHITE)
        c.drawString(MARGIN, PAGE_H - 76, title)
    else:
        c.setFillColor(NAVY)
        c.rect(0, PAGE_H - 88, PAGE_W, 88, fill=1, stroke=0)
        c.setFillColor(ORANGE)
        c.rect(0, PAGE_H - 92, PAGE_W, 4, fill=1, stroke=0)
        draw_logo(c, logo_path, PAGE_W - 158, PAGE_H - 69, 116)
        set_font(c, "Inter-SemiBold", 8.5, HexColor("#B9C6D8"))
        c.drawString(MARGIN, PAGE_H - 31, section.upper())
        set_font(c, "Syne-Bold", 22, WHITE)
        c.drawString(MARGIN, PAGE_H - 61, title)


def draw_footer(c, page_no, label="Align HCM | Internal use"):
    set_font(c, "Inter-Medium", 7.5, MUTED)
    c.drawString(MARGIN, 18, label)
    c.drawRightString(PAGE_W - MARGIN, 18, f"{page_no:02d}")


def draw_stat(c, x, y, w, value, label, accent=ORANGE, dark=True):
    fill = Color(1, 1, 1, alpha=0.08) if dark else WHITE
    stroke = Color(1, 1, 1, alpha=0.13) if dark else HAIRLINE
    round_rect(c, x, y, w, 76, fill, 14, stroke)
    c.setFillColor(accent)
    c.roundRect(x + 14, y + 14, 4, 48, 2, fill=1, stroke=0)
    set_font(c, "Syne-Bold", 24, WHITE if dark else NAVY)
    c.drawString(x + 30, y + 39, value)
    set_font(c, "Inter-Medium", 8.5, HexColor("#B9C6D8") if dark else MUTED)
    c.drawString(x + 30, y + 19, label.upper())


def draw_cover(c, logo_path, totals):
    c.setFillColor(NAVY)
    c.rect(0, 0, PAGE_W, PAGE_H, fill=1, stroke=0)
    c.setFillColor(Color(1, 0.42, 0.17, alpha=0.18))
    c.circle(PAGE_W - 30, PAGE_H - 15, 185, fill=1, stroke=0)
    c.setFillColor(Color(0.07, 0.43, 0.38, alpha=0.18))
    c.circle(PAGE_W - 95, 28, 136, fill=1, stroke=0)
    c.setFillColor(ORANGE)
    c.rect(0, 0, 10, PAGE_H, fill=1, stroke=0)
    draw_logo(c, logo_path, PAGE_W - 242, PAGE_H - 120, 188)
    set_font(c, "Inter-SemiBold", 9.5, HexColor("#B9C6D8"))
    c.drawString(54, PAGE_H - 70, "INTERNAL EDITORIAL BRIEF | PREPARED 10 AUG 2026")
    set_font(c, "Syne-Bold", 42, WHITE)
    c.drawString(54, PAGE_H - 140, "Top 30 Story")
    c.drawString(54, PAGE_H - 188, "Shortlist")
    c.setFillColor(ORANGE)
    c.roundRect(54, PAGE_H - 211, 224, 6, 3, fill=1, stroke=0)
    draw_wrapped(
        c,
        "A decision-ready portfolio of reference candidates, narrative readiness, representation gaps, and next validation moves.",
        54,
        PAGE_H - 242,
        500,
        "Inter-Medium",
        13,
        18,
        HexColor("#D8E0EA"),
        3,
    )
    stat_y = 78
    stat_w = (PAGE_W - 108 - 36) / 4
    stats = [
        (f"${totals['value']/1_000_000:.2f}M", "True USD"),
        (f"{totals['narratives']}/30", "Narratives ready"),
        (f"{totals['quotes']}/30", "Raven quotes"),
        (str(totals["industries"]), "Industries"),
    ]
    for i, (value, label) in enumerate(stats):
        draw_stat(c, 54 + i * (stat_w + 12), stat_y, stat_w, value, label, ORANGE if i < 3 else TEAL)
    c.showPage()


def draw_executive(c, logo_path, stories, totals, page_no):
    draw_header(c, "Executive readout", "Portfolio signal", logo_path)
    usable = totals["narratives"]
    gaps = 30 - usable
    missing_quotes = 30 - totals["quotes"]
    y = PAGE_H - 124
    cards = [
        ("20 stories have narrative coverage", f"{usable} of 30 already have drafted narrative content. The writing burden is concentrated in {gaps} stories, not the entire shortlist.", TEAL, TEAL_TINT),
        ("14 quotes still need sourcing", f"Raven supports {totals['quotes']} stories. Validation and quote collection remain the main public-use constraint.", ORANGE, RUST_TINT),
        ("Reference clearance is the gate", "Reference Status, Validator, and Validation Notes are still blank. No client should be treated as cleared for public naming yet.", BLUE, BLUE_TINT),
    ]
    card_w = (PAGE_W - 2 * MARGIN - 24) / 3
    for i, (title, body, accent, fill) in enumerate(cards):
        x = MARGIN + i * (card_w + 12)
        round_rect(c, x, y - 118, card_w, 118, fill, 14)
        c.setFillColor(accent)
        c.roundRect(x + 14, y - 28, 38, 4, 2, fill=1, stroke=0)
        draw_wrapped(c, title, x + 14, y - 48, card_w - 28, "Inter-Bold", 11, 14, NAVY, 2)
        draw_wrapped(c, body, x + 14, y - 80, card_w - 28, "Inter", 8.2, 11, SLATE, 4)

    top_values = sorted(stories, key=lambda s: float(s["Value (true USD)"] or 0), reverse=True)[:6]
    chart_y = 84
    chart_h = 190
    chart_x = MARGIN
    chart_w = 465
    round_rect(c, chart_x, chart_y, chart_w, chart_h, WHITE, 14, HAIRLINE)
    set_font(c, "Inter-Bold", 11, NAVY)
    c.drawString(chart_x + 16, chart_y + chart_h - 25, "Highest-value shortlisted stories")
    max_value = max(float(s["Value (true USD)"] or 0) for s in top_values)
    for i, story in enumerate(top_values):
        yy = chart_y + chart_h - 52 - i * 23
        name = ascii_text(story["Client"])
        value = float(story["Value (true USD)"] or 0)
        set_font(c, "Inter-Medium", 7.8, SLATE)
        c.drawString(chart_x + 16, yy + 2, name[:26])
        bar_x = chart_x + 156
        bar_w = 245 * (value / max_value)
        c.setFillColor(PANEL)
        c.roundRect(bar_x, yy, 245, 9, 4.5, fill=1, stroke=0)
        c.setFillColor(ORANGE if i < 2 else TEAL)
        c.roundRect(bar_x, yy, bar_w, 9, 4.5, fill=1, stroke=0)
        set_font(c, "Inter-Bold", 7.8, NAVY)
        c.drawRightString(chart_x + chart_w - 16, yy + 2, f"${value:,.0f}")

    index_x = chart_x + chart_w + 14
    index_w = PAGE_W - MARGIN - index_x
    round_rect(c, index_x, chart_y, index_w, chart_h, NAVY, 14)
    set_font(c, "Inter-Bold", 11, WHITE)
    c.drawString(index_x + 16, chart_y + chart_h - 25, "30-client identity index")
    set_font(c, "Inter", 7.2, HexColor("#B9C6D8"))
    c.drawString(index_x + 16, chart_y + chart_h - 41, "Embedded text lockups keep the PDF readable offline.")
    col_w = (index_w - 32) / 2
    for i, story in enumerate(stories):
        col = i // 15
        row = i % 15
        xx = index_x + 16 + col * col_w
        yy = chart_y + chart_h - 63 - row * 9.2
        set_font(c, "Inter-SemiBold", 6.2, WHITE)
        c.drawString(xx, yy, f"{int(story['Rank']):02d}  {ascii_text(story['Client'])[:28]}")
    draw_footer(c, page_no)
    c.showPage()


def draw_logo_wall(c, logo_path, stories, page_no):
    draw_header(c, "Client portfolio", "30 verified client identities", logo_path, dark=True)
    set_font(c, "Inter-Medium", 9.2, HexColor("#B9C6D8"))
    c.drawString(MARGIN, PAGE_H - 103, "Official marks are embedded in the PDF and repeated in each ranked story card.")
    cols, rows = 5, 6
    gap_x, gap_y = 10, 10
    chip_w = (PAGE_W - 2 * MARGIN - gap_x * (cols - 1)) / cols
    chip_h = 63
    start_y = PAGE_H - 132 - chip_h
    for idx, story in enumerate(stories):
        col, row = idx % cols, idx // cols
        x = MARGIN + col * (chip_w + gap_x)
        y = start_y - row * (chip_h + gap_y)
        round_rect(c, x, y, chip_w, chip_h, WHITE, 12)
        draw_client_logo(c, story, x + 9, y + 22, chip_w - 18, 34, False)
        set_font(c, "Inter-SemiBold", 6.7, MUTED)
        name = f"{int(story['Rank']):02d}  {ascii_text(story['Client'])}"
        c.drawCentredString(x + chip_w / 2, y + 8, name[:34])
    draw_footer(c, page_no, "Align HCM | Client portfolio | Internal use")
    c.showPage()


def status_chip(c, x, y, label, positive):
    fill = TEAL_TINT if positive else RUST_TINT
    color = TEAL if positive else RUST
    width = pdfmetrics.stringWidth(label, "Inter-SemiBold", 7.2) + 18
    round_rect(c, x, y, width, 18, fill, 9)
    set_font(c, "Inter-SemiBold", 7.2, color)
    c.drawCentredString(x + width / 2, y + 5.2, label)
    return width


def draw_story_card(c, story, x, y, w, h):
    round_rect(c, x, y, w, h, WHITE, 14, HAIRLINE)
    rank = int(story["Rank"])
    c.setFillColor(ORANGE)
    c.roundRect(x + 14, y + h - 45, 33, 33, 9, fill=1, stroke=0)
    set_font(c, "Syne-Bold", 15, WHITE)
    c.drawCentredString(x + 30.5, y + h - 34.5, f"{rank:02d}")

    name = ascii_text(story["Client"])
    draw_client_logo(c, story, x + 56, y + h - 45, 76, 34)
    set_font(c, "Syne-Bold", 13.5, NAVY)
    c.drawString(x + 143, y + h - 24, name)
    set_font(c, "Inter-Medium", 7.3, MUTED)
    employee_note = f" | {int(story['Employees (CRM)']):,} CRM employees" if story["Employees (CRM)"] else ""
    identity = f"{ascii_text(story['Website']) or 'Domain unavailable'} | {ascii_text(story['HQ (CRM)']) or ascii_text(story['Geo'])}{employee_note}"
    c.drawString(x + 143, y + h - 39, identity[:90])

    value = float(story["Value (true USD)"] or 0)
    set_font(c, "Syne-Bold", 15, NAVY)
    c.drawRightString(x + w - 16, y + h - 25, f"${value:,.0f}")
    set_font(c, "Inter-Medium", 6.8, MUTED)
    c.drawRightString(x + w - 16, y + h - 39, "TRUE USD")

    divider_y = y + h - 50
    c.setStrokeColor(HAIRLINE)
    c.setLineWidth(0.8)
    c.line(x + 14, divider_y, x + w - 14, divider_y)

    left_x = x + 16
    mid_x = x + 302
    right_x = x + w - 260
    label_y = divider_y - 13
    data_y = label_y - 13
    fields = [
        (left_x, "STORY", f"{ascii_text(story['Story ID'])} | {ascii_text(story['Story Type'])} | {int(story['Start Yr'])}"),
        (mid_x, "PORTFOLIO", f"{ascii_text(story['Industry'])} | {ascii_text(story['Size (HC)'])} | {ascii_text(story['Platform'])}"),
        (right_x, "CRM OWNER", ascii_text(story["CRM Deal Owner"])),
    ]
    widths = [270, 290, 240]
    for (xx, label, value_text), width in zip(fields, widths):
        set_font(c, "Inter-SemiBold", 6.6, MUTED)
        c.drawString(xx, label_y, label)
        draw_wrapped(c, value_text, xx, data_y, width, "Inter-SemiBold", 8.7, 10, SLATE, 1)

    chips_y = y + 5
    narrative_ok = ascii_text(story["Narrative?"]) == "Yes"
    quote_ok = ascii_text(story["Quote?"]) == "Raven"
    cx = left_x
    cx += status_chip(c, cx, chips_y, "Narrative ready" if narrative_ok else "Narrative needed", narrative_ok) + 7
    cx += status_chip(c, cx, chips_y, "Raven quote" if quote_ok else "Quote needed", quote_ok) + 7
    ref_status = ascii_text(story["Reference Status"]) or "Reference status pending"
    status_chip(c, cx, chips_y, ref_status, bool(story["Reference Status"]))
    set_font(c, "Inter-Medium", 7.1, RUST if story["What it still needs"] else MUTED)
    c.drawRightString(x + w - 16, chips_y + 5.5, f"Next: {ascii_text(story['What it still needs']) or 'Validation'}")


def draw_story_pages(c, logo_path, stories, start_page):
    page_no = start_page
    per_page = 4
    for start in range(0, len(stories), per_page):
        end = min(start + per_page, len(stories))
        subset = stories[start:end]
        draw_header(c, f"Stories {start + 1:02d}-{end:02d}", "Top 30 shortlist", logo_path)
        card_h = 101
        gap = 9
        y = PAGE_H - 112 - card_h
        for story in subset:
            draw_story_card(c, story, MARGIN, y, PAGE_W - 2 * MARGIN, card_h)
            y -= card_h + gap
        if len(subset) < per_page:
            panel_y = 54
            panel_h = 100
            round_rect(c, MARGIN, panel_y, PAGE_W - 2 * MARGIN, panel_h, NAVY, 14)
            set_font(c, "Syne-Bold", 14, WHITE)
            c.drawString(MARGIN + 18, panel_y + 73, "Close the validation loop")
            items = [
                ("Reference Status", "Use the approved ladder from None through Logo+Quote+Reference Call OK."),
                ("Validator", "Assign the delivery or CRM owner who can confirm the story and public-use scope."),
                ("Validation Notes", "Record consent, approved wording, quote source, and any logo-use limitations."),
            ]
            item_w = (PAGE_W - 2 * MARGIN - 36 - 24) / 3
            for idx, (title, body) in enumerate(items):
                xx = MARGIN + 18 + idx * (item_w + 12)
                set_font(c, "Inter-Bold", 8.5, ORANGE if idx != 1 else HexColor("#6ED6C6"))
                c.drawString(xx, panel_y + 49, title.upper())
                draw_wrapped(c, body, xx, panel_y + 33, item_w, "Inter", 7.5, 10, HexColor("#D8E0EA"), 3)
        draw_footer(c, page_no, "Align HCM | Top 30 shortlist | Internal use")
        c.showPage()
        page_no += 1
    return page_no


def table_text(c, value, x, y, width, font="Inter", size=7.3, color=SLATE, align="left", max_lines=2):
    text = ascii_text(value)
    lines = wrap_lines(text, font, size, width, max_lines)
    set_font(c, font, size, color)
    for i, line in enumerate(lines):
        yy = y - i * (size + 2)
        if align == "right":
            c.drawRightString(x + width, yy, line)
        elif align == "center":
            c.drawCentredString(x + width / 2, yy, line)
        else:
            c.drawString(x, yy, line)


def draw_reference_pages(c, logo_path, references, start_page):
    page_no = start_page
    widths = [102, 148, 120, 84, 42, 55, 88, 52, 84]
    for start in range(0, len(references), 10):
        subset = references[start:start + 10]
        draw_header(c, "Reference candidates", f"Objective 2 | {start + 1}-{start + len(subset)} of {len(references)}", logo_path)
        draw_wrapped(c, "Tier A already has a public Raven review. These remain candidates, not cleared public references.", MARGIN, PAGE_H - 112, PAGE_W - 2 * MARGIN, "Inter-Medium", 9, 12, MUTED, 2)
        x0 = MARGIN
        table_y = PAGE_H - 148
        header_h = 31
        c.setFillColor(NAVY)
        c.roundRect(x0, table_y - header_h, sum(widths), header_h, 10, fill=1, stroke=0)
        headers = ["Tier", "Client", "Industry", "Size", "Geo", "Eng.", "True USD", "Year", "Status"]
        xx = x0
        for title, width in zip(headers, widths):
            table_text(c, title.upper(), xx + 7, table_y - 19, width - 14, "Inter-SemiBold", 6.5, WHITE)
            xx += width
        y = table_y - header_h - 4
        row_h = 37
        for idx, row in enumerate(subset):
            fill = WHITE if idx % 2 == 0 else PANEL
            round_rect(c, x0, y - row_h, sum(widths), row_h, fill, 4)
            tier = ascii_text(row["Tier"])
            color = TEAL if tier.startswith("A") else BLUE
            xx = x0
            values = [
                tier,
                row["Client"],
                row["Industry"],
                row["Size (HC)"],
                row["Geo"],
                row["Engagements"],
                f"${float(row['Value (true USD)'] or 0):,.0f}",
                row["Latest yr"],
                row["Reference Status"] or "Pending",
            ]
            aligns = ["left", "left", "left", "left", "center", "center", "right", "center", "left"]
            for col, (value, width, align) in enumerate(zip(values, widths, aligns)):
                f = "Inter-SemiBold" if col in {0, 1, 6} else "Inter"
                col_color = color if col == 0 else (RUST if col == 8 else SLATE)
                table_text(c, value, xx + 7, y - 15, width - 14, f, 7.2, col_color, align, 2)
                xx += width
            y -= row_h + 3
        draw_footer(c, page_no, "Align HCM | Reference candidates | Internal use")
        c.showPage()
        page_no += 1
    return page_no


def draw_gap_section(c, section, x, top_y, w, max_rows=None):
    title = section["title"]
    rows = section["rows"][:max_rows] if max_rows else section["rows"]
    row_h = 21
    header_h = 24
    total_h = 35 + header_h + len(rows) * row_h + 4
    round_rect(c, x, top_y - total_h, w, total_h, WHITE, 12, HAIRLINE)
    set_font(c, "Inter-Bold", 10, NAVY)
    c.drawString(x + 12, top_y - 21, title)
    y = top_y - 35
    widths = [w * 0.35, w * 0.14, w * 0.14, w * 0.20, w * 0.17]
    c.setFillColor(NAVY)
    c.rect(x, y - header_h, w, header_h, fill=1, stroke=0)
    headers = ["Value", "Shortlist", "Book", "Book USD", "Flag"]
    xx = x
    for header, width in zip(headers, widths):
        table_text(c, header.upper(), xx + 6, y - 15, width - 12, "Inter-SemiBold", 6.2, WHITE)
        xx += width
    y -= header_h
    for idx, row in enumerate(rows):
        raw_flag = ascii_text(row[4])
        flag = "GAP" if raw_flag.startswith("GAP") else ("LOW VALUE" if raw_flag else "")
        fill = RUST_TINT if flag else (PANEL if idx % 2 else WHITE)
        c.setFillColor(fill)
        c.rect(x, y - row_h, w, row_h, fill=1, stroke=0)
        values = [row[0], row[1], row[2], f"${float(row[3] or 0):,.0f}", flag]
        aligns = ["left", "center", "center", "right", "left"]
        xx = x
        for col, (value, width, align) in enumerate(zip(values, widths, aligns)):
            color = RUST if flag and col == 4 else SLATE
            font = "Inter-SemiBold" if col in {0, 4} else "Inter"
            table_text(c, value, xx + 6, y - 14, width - 12, font, 6.4, color, align, 1)
            xx += width
        y -= row_h
    return total_h


def draw_gap_pages(c, logo_path, sections, start_page):
    page_no = start_page
    first = sections[:2]
    draw_header(c, "Representation gaps", "Objective 4 | Coverage by dimension", logo_path)
    col_w = (PAGE_W - 2 * MARGIN - 14) / 2
    draw_gap_section(c, first[0], MARGIN, PAGE_H - 112, col_w)
    draw_gap_section(c, first[1], MARGIN + col_w + 14, PAGE_H - 112, col_w)
    draw_footer(c, page_no, "Align HCM | Representation gaps | Internal use")
    c.showPage()
    page_no += 1

    draw_header(c, "Platform and geography", "Objective 4 | Coverage by dimension", logo_path)
    draw_gap_section(c, sections[2], MARGIN, PAGE_H - 132, col_w)
    draw_gap_section(c, sections[3], MARGIN + col_w + 14, PAGE_H - 132, col_w)
    round_rect(c, MARGIN, 84, PAGE_W - 2 * MARGIN, 108, NAVY, 14)
    set_font(c, "Syne-Bold", 15, WHITE)
    c.drawString(MARGIN + 18, 158, "Coverage signal")
    draw_wrapped(
        c,
        "The shortlist includes UKG Pro, UKG Ready, and Dayforce, with representation across US, Canada, multi-country, and other geographies. The explicit gaps are concentrated in story type and construction, not platform or geography.",
        MARGIN + 18,
        134,
        PAGE_W - 2 * MARGIN - 36,
        "Inter-Medium",
        10,
        14,
        HexColor("#D8E0EA"),
        3,
    )
    draw_footer(c, page_no, "Align HCM | Representation gaps | Internal use")
    c.showPage()
    page_no += 1

    story_type = next(s for s in sections if s["title"] == "STORY TYPE")
    chunks = [story_type["rows"][i:i + 13] for i in range(0, len(story_type["rows"]), 13)]
    for idx, chunk in enumerate(chunks):
        draw_header(c, "Story-type coverage", f"Objective 4 | Part {idx + 1} of {len(chunks)}", logo_path)
        pseudo = {"title": "STORY TYPE", "rows": chunk}
        draw_gap_section(c, pseudo, MARGIN, PAGE_H - 112, PAGE_W - 2 * MARGIN)
        draw_footer(c, page_no, "Align HCM | Representation gaps | Internal use")
        c.showPage()
        page_no += 1
    return page_no


def section_height(section, width):
    lines = 1
    for p in section["paragraphs"]:
        lines += len(wrap_lines(p, "Inter", 8.6, width - 28)) + 1
    return 42 + lines * 11


def draw_method_pages(c, logo_path, sections, start_page):
    page_no = start_page
    groups = [sections[:2], sections[2:4], sections[4:]]
    subtitles = ["Source, scope, and currency correction", "Selection logic and validation work", "Unblocking coverage and brand provenance"]
    for group, subtitle in zip(groups, subtitles):
        draw_header(c, "Method and operating notes", subtitle, logo_path)
        y = PAGE_H - 116
        for section in group:
            h = section_height(section, PAGE_W - 2 * MARGIN)
            h = min(h, y - 42)
            round_rect(c, MARGIN, y - h, PAGE_W - 2 * MARGIN, h, WHITE, 14, HAIRLINE)
            c.setFillColor(ORANGE)
            c.roundRect(MARGIN + 15, y - 28, 46, 4, 2, fill=1, stroke=0)
            set_font(c, "Syne-Bold", 13, NAVY)
            c.drawString(MARGIN + 15, y - 48, section["title"])
            text_y = y - 68
            for para in section["paragraphs"]:
                text_y = draw_wrapped(c, para, MARGIN + 15, text_y, PAGE_W - 2 * MARGIN - 30, "Inter", 8.6, 11.2, SLATE)
                text_y -= 7
            y -= h + 12
        draw_footer(c, page_no, "Align HCM | Method notes | Internal use")
        c.showPage()
        page_no += 1
    return page_no


def build_pdf():
    register_fonts()
    logo_path = extract_align_logo()
    stories, totals, references, sections, method_sections = load_data()
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(OUTPUT), pagesize=(PAGE_W, PAGE_H), pageCompression=1)
    c.setTitle("Align HCM Top 30 Story Shortlist - Editorial PDF")
    c.setAuthor("Align HCM")
    c.setSubject("Top 30 story shortlist, reference candidates, representation gaps, and method notes")
    c.setCreator("Align HCM branded deliverable workflow")

    draw_cover(c, logo_path, totals)
    page_no = 2
    draw_executive(c, logo_path, stories, totals, page_no)
    page_no += 1
    draw_logo_wall(c, logo_path, stories, page_no)
    page_no += 1
    page_no = draw_story_pages(c, logo_path, stories, page_no)
    page_no = draw_reference_pages(c, logo_path, references, page_no)
    page_no = draw_gap_pages(c, logo_path, sections, page_no)
    page_no = draw_method_pages(c, logo_path, method_sections, page_no)
    c.save()
    return OUTPUT


if __name__ == "__main__":
    output = build_pdf()
    print(output)
