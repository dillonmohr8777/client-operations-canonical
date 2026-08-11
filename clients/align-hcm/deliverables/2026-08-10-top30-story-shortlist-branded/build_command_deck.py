#!/usr/bin/env python3
"""Align HCM · Top 30 Command Deck v3 — server-rendered, mobile-first.

v3: every section is rendered into the HTML at build time, so the deck is
fully readable with JavaScript disabled or sandboxed (mobile file previews,
strict webviews). JS only upgrades: tabs, filters, sort, search, drawer,
tooltips, motion, localStorage state. Plus a full responsive pass: scaled
constellation, compact header, tightened grid tracks, tap-friendly controls,
100dvh drawer. Data and brand identical to v2 (HubSpot portal 242825734 read
10 Aug 2026; won/open/lost split by real deal state).
"""
import base64, json, math, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = "/root/.claude/uploads/9e045f79-a3fe-5e0b-b800-e5900182d04c/fc733806-Align_HCM_Top30_Story_Shortlist.xlsx"
LOGO = os.path.join(HERE, "logos/align-hcm-logo-reverse.png")
OUT = os.path.join(HERE, "build/Align_HCM_Top30_Command_Deck.html")
STAMP = "10 Aug 2026 · 22:10 UTC"
PORTAL = "242825734"
HS_BASE = f"https://app-na2.hubspot.com/contacts/{PORTAL}/company/"

# ------------------------------------------------------------------ data ----
wb = openpyxl.load_workbook(SRC, data_only=True)
s30, sref, sgap = (wb[n] for n in ("Top 30 Shortlist","Reference Candidates","Representation Gaps"))
def v(ws,r,c):
    x = ws.cell(row=r,column=c).value
    return "" if x is None else x
STORIES = [dict(rank=int(v(s30,r,1)), sid=v(s30,r,2), client=v(s30,r,3), industry=v(s30,r,4),
    size=str(v(s30,r,5)), geo=v(s30,r,6), platform=v(s30,r,7), stype=v(s30,r,8),
    yr=int(v(s30,r,9)), value=round(float(v(s30,r,10))), narr=v(s30,r,11),
    quote=v(s30,r,12), needs=v(s30,r,13), owner=v(s30,r,14)) for r in range(6,36)]
REFS = [dict(tier=v(sref,r,1), client=v(sref,r,2), industry=v(sref,r,3), size=str(v(sref,r,4)),
    geo=v(sref,r,5), eng=int(v(sref,r,6)), value=round(float(v(sref,r,7))), yr=int(v(sref,r,8)))
    for r in range(5, sref.max_row+1) if v(sref,r,2)]
GAPS, cur = [], None
for r in range(3, sgap.max_row+1):
    a = str(v(sgap,r,1))
    if not any(str(v(sgap,r,c)) != "" for c in range(1,6)): continue
    if a.isupper() and v(sgap,r,2) == "":
        cur = dict(dim=a, rows=[]); GAPS.append(cur)
    elif a == "Value" or cur is None: continue
    else: cur["rows"].append([a, int(v(sgap,r,2)), int(v(sgap,r,3)), round(float(v(sgap,r,4))), str(v(sgap,r,5)).strip()])

ENRICH = {
 "Leyad": ("leyad.ca","leyad.ca","Ontario, Canada",300),
 "GES": ("ges.com","ges.com","Las Vegas, NV",""),
 "Curtin Maritime": ("curtinmaritime.com","curtinmaritime.com","",500),
 "EverQuote": ("everquote.com","everquote.com","",350),
 "BURNCO": ("burnco.com","burnco.com","Calgary, AB",1500),
 "OhioHealth": ("ohiohealth.com","ohiohealthems.com","Columbus, OH",30000),
 "McLeod Software": ("mcleodsoftware.com","mcleodsoftware.com","Birmingham, AL",700),
 "PrimeFlight": ("primeflight.com","primeflight.com","Sugar Land, TX",9400),
 "Vacasa": ("vacasa.com","vacasa.com","Portland, OR",6800),
 "MW Components": ("mwcomponents.com","mwcomponents.com","Charlotte, NC",2100),
 "Distributor Wire & Cable": ("distributorwireandcable.com","distributorwireandcable.com","",""),
 "Bethany Christian Services": ("bethany.org","bethany.org","",""),
 "TriState Health": ("tsh.org","tsh.org","Clarkson, WA",""),
 "KW Property Management": ("kwpmc.com","kwpmc.com","Miami, FL",2700),
 "Troon": ("troon.com","troon.com","Scottsdale, AZ",30000),
 "JACAM Catalyst": ("jacamcatalyst.com",None,"Sterling, KS",""),
 "Chautauqua Institution": ("chq.org","chq.org","",1400),
 "Excelsior University": ("excelsior.edu","excelsior.edu","Albany, NY",""),
 "Eagle Materials": ("eaglematerials.com","eaglematerials.com","Dallas, TX",2657),
 "MTA NY": ("mta.info","mtany.org","",70000),
 "Grand River Health": ("grandriverhealth.org","grandriverhealth.org","Rifle, CO",850),
 "Bayshore HealthCare": ("bayshore.ca","bayshore.ca","Etobicoke, ON",3500),
 "Rollins": ("rollins.com","rollins.com","Atlanta, GA",13000),
 "Driscoll's": ("driscolls.com","driscolls.com","Watsonville, CA",6200),
 "Trimac Transportation": ("trimac.com","trimac.com","Calgary, AB",""),
 "Coastal Community Credit Union": ("cccu.ca","cccu.ca","",""),
 "REI": ("rei.com","rei.com","",""),
 "UBE": ("ube.com","ube.com","",""),
 "Ace Parking": ("aceparking.com",None,"",5000),
 "McCain Foods": ("mccain.com","mccain.ca","Ontario, Canada",20000),
}

ROLLUP = {d: dict(crm_name=nm, crm_id=oid, deals=int(n), lifetime=round(float(amt)))
  for d, nm, oid, n, amt in (l.split("|") for l in """bayshore.ca|bayshore healthcare|123711651550|4|555511.41
bethany.org|bethany christian services|123255253734|5|186095
burnco.com|burnco rock products ltd.|123183723207|13|567352.88
cccu.ca|coastal community credit union|314030019312|1|213815.30
chq.org|chautauqua institution|123255253744|5|126875
curtinmaritime.com|curtin maritime, corp.|217755393775|2|168000
distributorwireandcable.com|distributor wire & cable|123181851364|1|54000
driscolls.com|driscoll's inc|123181851369|22|1274698.01
eaglematerials.com|eagle materials inc.|123155562231|20|741000
everquote.com|everquote|123180069576|3|119425
excelsior.edu|excelsior university|123185434320|2|79000
ges.com|ges|123155562214|4|279500
grandriverhealth.org|grand river health|123255253725|5|426650
kwpmc.com|kw property management & consulting|123713450713|3|230500
leyad.ca|leyad|326104277736|2|76112.14
mccain.ca|mccain foods|231724321469|1|83426.60
mcleodsoftware.com|mcleod software|123672827600|1|96875
mtany.org|mta ny|123152084691|4|5296000
mwcomponents.com|mw components|123150291674|4|290550
ohiohealthems.com|ohiohealth corporation|320024648395|4|2932695
primeflight.com|primeflight aviation services, inc.|123181851383|14|1984250
rei.com|recreational equipment inc.|271741020870|4|990950
rollins.com|rollins, inc.|123713450709|4|52750
trimac.com|trimac transportation|123726965453|3|475656.52
troon.com|troon golf, llc|123155562216|14|564500
tsh.org|tristate health|123227122411|1|142950
ube.com|ube corporation americas inc|219952394992|2|229900
vacasa.com|vacasa llc|123181851374|2|360000""".splitlines())}
NOTWON = {}
for part in """bayshore.ca|1|97221.82;bethany.org|1|18000;burnco.com|1|3294.41;chq.org|1|11000;curtinmaritime.com|1|84000;driscolls.com|4|345970;eaglematerials.com|1|10000;everquote.com|2|40625;ges.com|2|134000;kwpmc.com|1|52500;mtany.org|1|4000000;mwcomponents.com|2|35000;primeflight.com|3|566400;rei.com|2|611000;troon.com|4|115000;ube.com|1|73600;vacasa.com|1|120000""".split(";"):
    d, n, amt = part.split("|"); NOTWON[d] = (int(n), round(float(amt)))
WON = {}
for part in """bayshore.ca|3|458289.59;bethany.org|4|168095;burnco.com|12|564058.46;cccu.ca|1|213815.30;chq.org|4|115875;curtinmaritime.com|1|84000;distributorwireandcable.com|1|54000;driscolls.com|18|928728.01;eaglematerials.com|19|731000;everquote.com|1|78800;excelsior.edu|2|79000;ges.com|2|145500;grandriverhealth.org|5|426650;kwpmc.com|2|178000;leyad.ca|2|76112.14;mccain.ca|1|83426.60;mcleodsoftware.com|1|96875;mtany.org|3|1296000;mwcomponents.com|2|255550;ohiohealthems.com|4|2932695;primeflight.com|11|1417850;rei.com|2|379950;rollins.com|4|52750;trimac.com|3|475656.52;troon.com|10|449500;tsh.org|1|142950;ube.com|1|156300;vacasa.com|1|240000""".split(";"):
    d, n, amt = part.split("|"); WON[d] = (int(n), round(float(amt)))
OPEN_DEALS = [
 dict(dom="primeflight.com", name="PrimeFlight - UK Full-suite Launch", amt=500000, stage="Expressing Interest", close="2026-11-30", owner="Allison Cox"),
 dict(dom="rei.com", name="Recreational Equipment Inc. - PRO/WFM Business Structure Alignment", amt=411000, stage="Qualification", close="2026-09-30", owner="Michael Lederman"),
 dict(dom="rei.com", name="Recreational Equipment Inc, SmartCare", amt=200000, stage="Qualification", close="2026-09-30", owner="Michael Lederman"),
 dict(dom="curtinmaritime.com", name="Curtin Maritime SmartCare - Renewal", amt=84000, stage="Expressing Interest", close="2026-11-30", owner="Michael Lederman"),
 dict(dom="ube.com", name="UBE - CR028-2026", amt=73600, stage="Contracting", close="2026-08-31", owner="Dianna Hammond"),
 dict(dom="everquote.com", name="EverQuote - CR061-2026", amt=34875, stage="Contracting", close="2026-08-21", owner="Dianna Hammond"),
]
YEARS_WON = {}
for part in """2024|bayshore.ca|3|458290;2024|bethany.org|1|153495;2025|bethany.org|2|13000;2026|bethany.org|1|1600;2024|burnco.com|1|108000;2025|burnco.com|8|193225;2026|burnco.com|3|262833;2026|cccu.ca|1|213815;2024|chq.org|1|78375;2025|chq.org|2|19500;2026|chq.org|1|18000;2025|curtinmaritime.com|1|84000;2025|distributorwireandcable.com|1|54000;2019|driscolls.com|1|8053;2020|driscolls.com|3|60500;2021|driscolls.com|4|54325;2022|driscolls.com|1|72000;2023|driscolls.com|3|493850;2024|driscolls.com|2|70000;2025|driscolls.com|2|110000;2026|driscolls.com|2|60000;2022|eaglematerials.com|1|120000;2023|eaglematerials.com|4|140000;2024|eaglematerials.com|8|268000;2025|eaglematerials.com|5|163000;2026|eaglematerials.com|1|40000;2025|everquote.com|1|78800;2024|excelsior.edu|1|75000;2025|excelsior.edu|1|4000;2025|ges.com|1|143000;2026|ges.com|1|2500;2025|grandriverhealth.org|1|317700;2026|grandriverhealth.org|4|108950;2024|kwpmc.com|1|138000;2025|kwpmc.com|1|40000;2026|leyad.ca|2|76112;2026|mccain.ca|1|83427;2025|mcleodsoftware.com|1|96875;2024|mtany.org|1|432000;2025|mtany.org|1|432000;2026|mtany.org|1|432000;2024|mwcomponents.com|1|231550;2025|mwcomponents.com|1|24000;2024|ohiohealthems.com|2|2165175;2025|ohiohealthems.com|1|460020;2026|ohiohealthems.com|1|307500;2023|primeflight.com|2|1090360;2024|primeflight.com|5|296490;2025|primeflight.com|4|31000;2026|rei.com|2|379950;2022|rollins.com|1|46500;2025|rollins.com|3|6250;2026|trimac.com|3|475657;2023|troon.com|2|288000;2024|troon.com|2|64000;2025|troon.com|5|97500;2026|troon.com|1|0;2024|tsh.org|1|142950;2026|ube.com|1|156300;2024|vacasa.com|1|240000""".split(";"):
    y, d, n, amt = part.split("|"); YEARS_WON.setdefault(d, {})[int(y)] = [int(n), round(float(amt))]
PULSE = [
  dict(name="UKG Sales Pipeline", won_n=755, won=53681656, open_n=711, open=100086750),
  dict(name="Dayforce Sales Pipeline", won_n=4, won=584487, open_n=44, open=14421746),
  dict(name="Paylocity Sales Pipeline", won_n=0, won=0, open_n=3, open=660000),
]
OWNERS_ACTIVE = {"Mike Emsley": True, "Maher El-Abdallah": True, "Dianna Hammond": True,
                 "Michael Lederman": True, "Miranda Sarwan": False, "Allison Cox": True}

open_by_dom = {}
for d_ in OPEN_DEALS:
    e = open_by_dom.setdefault(d_["dom"], [0,0]); e[0]+=1; e[1]+=d_["amt"]
for s in STORIES:
    pub, crm_dom, hq, emp = ENRICH[s["client"]]
    s.update(domain=pub, hq=hq, emp=emp, ownerActive=OWNERS_ACTIVE.get(s["owner"]), hs=None)
    if crm_dom and crm_dom in ROLLUP:
        r = ROLLUP[crm_dom]
        wn, wv = WON.get(crm_dom, (0,0)); nn, nv = NOTWON.get(crm_dom, (0,0))
        on, ov = open_by_dom.get(crm_dom, (0,0))
        s["hs"] = dict(crm_name=r["crm_name"], crm_id=r["crm_id"], deals=r["deals"],
            lifetime=r["lifetime"], won=wv, won_n=wn, open=ov, open_n=on,
            lost=nv-ov, lost_n=nn-on, years=YEARS_WON.get(crm_dom, {}),
            deals_open=[x for x in OPEN_DEALS if x["dom"]==crm_dom])
for rr in REFS:
    e = ENRICH.get(rr["client"]); rr["domain"] = e[0] if e else None

# ------------------------------------------------------- python renderers ---
def esc(x):
    return (str(x).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;").replace('"',"&quot;"))
fmt = lambda n: "$" + format(round(n), ",")
def fmtM(n):
    if n >= 1e6: return f"${n/1e6:.2f}M"
    if n >= 1e3: return f"${round(n/1e3)}k"
    return fmt(n)
def initials(name):
    parts = [w for w in str(name).replace("&"," ").split() if w and (w[0].isalnum())]
    return "".join(w[0] for w in parts[:2]).upper()
def readiness(s):
    return "ready" if s["needs"].startswith("Ready") else ("fresh" if s["needs"].startswith("Fresh") else "todo")
READY_LABEL = {"ready":"Ready · spot-check", "fresh":"Needs fresh quote", "todo":"Needs narrative + quote"}
READY_COLOR = {"ready":"#1FA98C", "fresh":"#4E86C8", "todo":"#F05A28"}
PLATC = {"UKG Pro":"c-pro","UKG Ready":"c-ready","Dayforce":"c-day"}
STAGEC = {"Expressing Interest":"st-int","Qualification":"st-qual","Contracting":"st-con"}
def logo_chip(s, sz):
    return (f'<span class="lg" style="--sz:{sz}px"><img loading="lazy" src="https://logo.clearbit.com/{s["domain"]}" alt="" '
            f'onerror="this.parentElement.classList.add(\'fb\')"><i>{initials(s["client"])}</i></span>')

bookVal = sum(s["value"] for s in STORIES)
withHS = [s for s in STORIES if s["hs"]]
openTot = sum(s["hs"]["open"] for s in withHS); openN = sum(s["hs"]["open_n"] for s in withHS)
wonTot = sum(s["hs"]["won"] for s in withHS); lostTot = sum(s["hs"]["lost"] for s in withHS)
inactive = [s for s in STORIES if s["ownerActive"] is False]
narrN = sum(1 for s in STORIES if s["narr"]=="Yes"); quoN = sum(1 for s in STORIES if s["quote"]=="Raven")

def kpi(lab, val, ctx="", cls="", ac=""):
    st = f' style="--ac:{ac}"' if ac else ""
    return (f'<div class="kpi {cls} reveal"{st}><span class="lab">{lab}</span>'
            f'<b class="num" data-final="{esc(val)}">{esc(val)}</b>'
            + (f'<div class="ctx">{ctx}</div>' if ctx else "") + '</div>')

KPIS = (kpi("Story book", fmtM(bookVal), "true USD · one story per client", "hot")
 + kpi("Narratives drafted", f"{narrN} of 30", "10 still to write")
 + kpi("Raven quotes", f"{quoN} of 30", "14 to source")
 + kpi("Lifetime won", fmtM(wonTot), "these 30 · closed-won all time", "", "#1FA98C")
 + kpi("Open now", fmtM(openTot), f"{openN} live deals", "", "#FF6B2B")
 + kpi("Closed lost", fmtM(lostTot), "incl. one $4.0M loss", "", "#5F7189")
 + kpi("Owner risk", f"{len(inactive)} stories", "owner inactive in CRM", "", "#E5484D"))

# constellation
rings = {"ready":(152,"95s",11), "fresh":(206,"120s",45), "todo":(258,"150s",18)}
C = 300
groups = {"ready":[],"fresh":[],"todo":[]}
for s in STORIES: groups[readiness(s)].append(s)
vmax = max(s["value"] for s in STORIES)
orbit = [f'<div class="coreglow"></div><div class="core"><img src="data:image/png;base64,@@LOGO@@" alt="Align HCM"><b>{fmtM(bookVal)}</b><span>Thirty stories · three rings</span></div>']
for key in ("ready","fresh","todo"):
    r, dur, ph = rings[key]
    orbit.append(f'<div class="ringpath" style="inset:{C-r}px;--rd:{int(dur[:-1])*2.4:.0f}s"></div><div class="ring" style="--dur:{dur}">')
    n = len(groups[key])
    for i, s in enumerate(groups[key]):
        ang = (i/n)*2*math.pi - math.pi/2 + ph*math.pi/180
        sz = round(28 + 24*math.sqrt(s["value"]/vmax))
        x, y = C + r*math.cos(ang), C + r*math.sin(ang)
        tipc = f"<b>{esc(s['client'])}</b><br>{fmt(s['value'])} · {READY_LABEL[key]}<br><span class=m>#{s['rank']} · {esc(s['platform'])} · tap for detail</span>"
        orbit.append(f'<div class="node" data-client="{esc(s["client"])}" data-tip="{esc(tipc)}" '
            f'style="--s:{sz}px;--hc:{READY_COLOR[key]};left:{x:.1f}px;top:{y:.1f}px">'
            f'<img loading="lazy" src="https://logo.clearbit.com/{s["domain"]}" alt="" onerror="this.parentElement.classList.add(\'fb\')"><i>{initials(s["client"])}</i></div>')
    orbit.append('</div>')
ORBIT = "".join(orbit)
ORBITLEGEND = "".join(f'<span class="olg"><span class="dot" style="background:{READY_COLOR[k]};box-shadow:0 0 8px {READY_COLOR[k]}"></span>{READY_LABEL[k]} · {len(groups[k])}</span>' for k in ("ready","fresh","todo"))

FEED = "".join(f'<div class="sig" data-go="{go}" style="--sc:{c}"><span class="dot"></span><div><span class="k">{k}</span><p>{t}</p></div><span class="go">›</span></div>'
  for c,k,go,t in [
    ("#FF6B2B","Money in motion","pipeline", f"<b>{fmtM(openTot)}</b> open across <b>{openN}</b> live deals. Largest: PrimeFlight UK full-suite launch, <b>$500k</b>, expressing interest, closes 30 Nov."),
    ("#1FA98C","2026 surge","pipeline", "<b>REI</b> is all-2026: <b>$379,950</b> won plus <b>$611k</b> open in qualification. Story #27 rides the wave."),
    ("#E5484D","Recompete signal","pipeline", "<b>MTA NY</b> lost a <b>$4.0M</b> deal in 2025, while $432k of staff-aug work still lands each year. Story #20 is the proof asset for the rematch."),
    ("#BA841A","Owner risk","ops", f"<b>{len(inactive)} stories</b> sit with an inactive CRM owner. Reassign before validation outreach."),
])

def brow(lab, pct, val, cls, tipc):
    return (f'<div class="brow reveal" data-tip="{esc(tipc)}"><span class="lab">{esc(lab)}</span>'
            f'<span class="track"><span class="fill {cls}" style="width:{pct:.1f}%"></span></span>'
            f'<span class="val">{val}</span></div>')
top7 = sorted(STORIES, key=lambda s:-s["value"])[:7]
TOPBARS = "".join(brow(s["client"], max(2,100*s["value"]/top7[0]["value"]), fmtM(s["value"]), "o",
    f"<b>{esc(s['client'])}</b><br>{fmt(s['value'])} story value · #{s['rank']}") for s in top7)
gcnt = {k:len(groups[k]) for k in groups}
READYBARS = "".join(brow(lab, 100*n/30, f'{n} <span style="color:var(--ink3)">/ 30</span>', cls,
    f"<b>{lab}</b><br>{n} of 30 stories") for lab,n,cls in
    [("Ready · spot-check only",gcnt["ready"],"t"),("Needs a fresh quote",gcnt["fresh"],"s"),("Needs narrative + quote",gcnt["todo"],"o")])

def card(s):
    r = readiness(s)
    owner_bit = (f'<span class="chip c-crit" data-tip="<b>{esc(s["owner"])}</b> is inactive in HubSpot.<br>Reassign before validation outreach.">⚠ {esc(s["owner"])}</span>'
        if s["ownerActive"] is False else f'<span class="chip c-mut">{esc(s["owner"])}</span>')
    if s["hs"]:
        h = s["hs"]; closed = h["won_n"]+h["lost_n"]
        win = f"{round(100*h['won_n']/closed)}%" if closed else "·"
        hs = (f'<div class="hsrow"><span data-tip="Closed-won lifetime, {h["won_n"]} deals">Won <b>{fmtM(h["won"])}</b></span>'
              f'<span class="open" data-tip="Open right now{" ("+str(h["open_n"])+" deals)" if h["open_n"] else ""}">Open <b>{fmtM(h["open"]) if h["open"] else "·"}</b></span>'
              f'<span data-tip="Win rate on closed deals">Win <b>{win}</b></span>'
              f'<span style="margin-left:auto"><span class="livechip">LIVE</span></span></div>')
    else:
        hs = '<div class="hsrow"><span style="color:var(--ink3)">No CRM domain link · story data only</span></div>'
    q = f'{s["client"]} {s["industry"]} {s["stype"]} {s["owner"]} {s["geo"]}'.lower()
    return f"""<article class="card reveal" data-client="{esc(s["client"])}" data-plat="{esc(s["platform"])}" data-ready="{r}"
 data-rank="{s['rank']}" data-value="{s['value']}" data-open="{s['hs']['open'] if s['hs'] else 0}" data-won="{s['hs']['won'] if s['hs'] else 0}" data-q="{esc(q)}">
<span class="rank {"top" if s["rank"]<=5 else ""}">#{s["rank"]}</span>
<div class="top">{logo_chip(s,46)}<div><h3>{esc(s["client"])}</h3><div class="sid">{esc(s["sid"])} · {esc(s["domain"])}</div></div></div>
<div class="money num">{fmt(s["value"])}</div>
<div class="vbar"><i style="width:{max(2,100*s["value"]/vmax):.1f}%"></i></div>
<div class="mline">{esc(s["industry"])} · {esc(s["size"])} HC · {esc(s["geo"])}{" · HQ "+esc(s["hq"]) if s["hq"] else ""}</div>
<div class="chips"><span class="chip {PLATC.get(s["platform"],"")}">{esc(s["platform"])}</span>
<span class="chip c-mut">{esc(s["stype"])}</span>
<span class="chip {"c-ok" if r=="ready" else "c-fresh" if r=="fresh" else "c-todo"}">{READY_LABEL[r]}</span>
{owner_bit}</div>{hs}</article>"""
CARDS = "".join(card(s) for s in STORIES)

PULSE_HTML = (kpi("Portal closed won · all time", fmtM(sum(p["won"] for p in PULSE)), f'{sum(p["won_n"] for p in PULSE)} deals', "", "#1FA98C")
 + kpi("Portal open pipeline", fmtM(sum(p["open"] for p in PULSE)), f'{sum(p["open_n"] for p in PULSE)} deals', "", "#FF6B2B")
 + "".join(kpi(p["name"], fmtM(p["won"]+p["open"]), f'{p["won_n"]} won · {p["open_n"]} open', "", "#4E86C8") for p in PULSE))

by_dom = {s["domain"]: s for s in STORIES}
deal_rows = []
for d in OPEN_DEALS:
    s = by_dom.get(d["dom"])
    chip = logo_chip(s,44) if s else '<span class="lg fb" style="--sz:44px"><i>?</i></span>'
    deal_rows.append(f'<div class="dealrow">{chip}<div><div class="dn">{esc(d["name"])}</div>'
        f'<div class="dm"><span class="stage {STAGEC.get(d["stage"],"")}">{esc(d["stage"])}</span>{esc(d["owner"])} · closes {d["close"]}</div></div>'
        f'<div class="da"><b class="num">{fmt(d["amt"])}</b><span>open</span></div></div>')
DEALSBOARD = "".join(deal_rows) + (f'<div class="dealsum"><span>{len(OPEN_DEALS)} open deals on shortlist clients</span>'
    f'<b class="num">{fmt(sum(d["amt"] for d in OPEN_DEALS))}</b></div>')

rows_wo = sorted(withHS, key=lambda s: -(s["hs"]["won"]+s["hs"]["open"]))
wmax = max(max(s["hs"]["won"], s["hs"]["open"]) for s in rows_wo)
wo_rows = []
for s in rows_wo:
    h = s["hs"]
    tipc = (f"<b>{esc(s['client'])}</b><br>Won lifetime: {fmt(h['won'])} ({h['won_n']})"
            f"<br>Open now: {fmt(h['open'])+' ('+str(h['open_n'])+')' if h['open'] else 'none'}"
            f"<br>Closed lost: {fmt(h['lost'])+' ('+str(h['lost_n'])+')' if h['lost'] else 'none'}")
    wo_rows.append(f'<div class="brow duo reveal" data-tip="{esc(tipc)}"><span class="lab">{esc(s["client"])}</span>'
        f'<span class="track"><span class="fill s a" style="width:{max(1.5,100*h["won"]/wmax):.1f}%"></span>'
        f'<span class="fill o b" style="width:{max(1.5,100*h["open"]/wmax) if h["open"] else 0:.1f}%"></span></span>'
        f'<span class="val">{fmtM(h["won"]+h["open"])}</span></div>')
WONOPEN = "".join(wo_rows)

ymax = max(v_[1] for s in withHS for v_ in s["hs"]["years"].values())
def shade(val):
    if val <= 0: return "rgba(240,90,40,.10)"
    return f"rgba(240,90,40,{0.16+0.84*math.sqrt(val/ymax):.2f})"
heat_rows = ['<div class="heatrow"><span class="hl"></span>' +
    "".join(f'<span class="heathead">{y}</span>' for y in range(2019,2027)) +
    '<span class="heathead" style="text-align:right">Won</span></div>']
for s in sorted(withHS, key=lambda s:-s["hs"]["won"]):
    cells = []
    for y in range(2019,2027):
        vv = s["hs"]["years"].get(y)
        tipc = f"<b>{esc(s['client'])} · {y}</b><br>" + (f"{vv[0]} deal{'s' if vv[0]>1 else ''} won · {fmt(vv[1])}" if vv else "No wins")
        style = f' style="background:{shade(vv[1])}"' if vv else ""
        cells.append(f'<span class="cell"{style} data-tip="{esc(tipc)}"></span>')
    heat_rows.append(f'<div class="heatrow"><span class="hl">{esc(s["client"])}</span>{"".join(cells)}<span class="heattot">{fmtM(s["hs"]["won"])}</span></div>')
HEAT = "".join(heat_rows)

ALERT = (f'<div><b>{len(inactive)} of 30 stories are owned by an inactive CRM user ({esc(inactive[0]["owner"])}).</b>'
    f'<p>Owner is inactive in HubSpot as of {STAMP}. Reassign these before validation outreach: '
    + " · ".join(esc(s["client"]) for s in inactive) + '</p></div>')

by_owner = {}
for s in STORIES: by_owner.setdefault(s["owner"], []).append(s)
owner_cards = []
for idx, (owner, lst) in enumerate(sorted(by_owner.items(), key=lambda kv:-len(kv[1]))):
    active = lst[0]["ownerActive"]; unassigned = "not in CRM" in owner
    val = sum(s["value"] for s in lst)
    badge = ('<span class="chip c-crit">inactive</span>' if active is False else
             '<span class="chip c-ok">active</span>' if active else "")
    items = "".join(f'<div class="oitem" data-sid="{esc(s["sid"])}" data-needs="{esc(s["needs"])}"><b>#{s["rank"]}</b> {esc(s["client"])}<span class="sp"></span>'
        f'<span class="chip {"c-fresh" if readiness(s)=="ready" else "c-todo"} ost-chip">{esc(s["needs"])}</span></div>' for s in lst)
    owner_cards.append(f'<div class="panel ocard reveal d{(idx%4)+1}">'
        f'<div class="oh"><span class="avatar {"gray" if (active is False or unassigned) else ""}">{"?" if unassigned else initials(owner)}</span>'
        f'<div><h3>{esc("Unassigned in CRM" if unassigned else owner)} {badge}</h3>'
        f'<div class="ost">{len(lst)} stories · {fmtM(val)} story value</div></div></div>'
        f'<div class="olist">{items}</div>'
        f'<button class="copybtn" data-owner="{esc(owner)}">Copy ask list</button></div>')
OWNERS_HTML = "".join(owner_cards)

cov_cards = []
for i, g in enumerate(GAPS):
    gmax = max(r_[3] for r_ in g["rows"])
    row_parts = []
    for lab, inS, inB, val, flag in g["rows"]:
        is_gap = flag.startswith("GAP")
        tipc = f"<b>{esc(lab)}</b><br>{inS} shortlisted of {inB} in book<br>Book value {fmt(val)}"
        if flag: tipc += "<br>" + esc(flag)
        flag_html = ""
        if flag:
            fcls = "c-todo" if is_gap else "c-warn"
            ftxt = "GAP" if is_gap else "low value"
            flag_html = f'<span class="chip {fcls}">{ftxt}</span>'
        row_parts.append(
            f'<div class="gaprow" data-tip="{esc(tipc)}">'
            f'<span class="lab">{esc(lab)}</span>'
            f'<span class="track"><span class="fill {"o" if is_gap else "s"}" style="width:{max(2,100*val/gmax):.1f}%"></span></span>'
            f'<span class="n">{inS} / {inB}</span>'
            f'<span>{flag_html}</span></div>')
    cov_cards.append(f'<div class="panel reveal d{(i%3)+1}"><span class="kick" style="color:#E8B95C">{esc(g["dim"])}</span><div class="chart">{"".join(row_parts)}</div></div>')
COV = "".join(cov_cards)

bench_rows = []
for r in REFS:
    s = next((x for x in STORIES if x["client"]==r["client"]), None)
    live = fmtM(s["hs"]["won"]) if s and s["hs"] else "·"
    bench_rows.append(f'<tr><td><span class="tier {"a" if r["tier"].startswith("A") else "b"}">{esc(r["tier"])}</span></td>'
        f'<td><b>{esc(r["client"])}</b></td><td>{esc(r["industry"])}</td><td>{esc(r["size"])}</td><td>{esc(r["geo"])}</td>'
        f'<td>{r["eng"]}</td><td><b class="num">{fmt(r["value"])}</b></td><td class="num">{live}</td><td>{r["yr"]}</td></tr>')
BENCH = ('<thead><tr><th>Tier</th><th>Client</th><th>Industry</th><th>Size</th><th>Geo</th><th>Engagements</th>'
         '<th>Sheet value</th><th>CRM won · live</th><th>Latest yr</th></tr></thead><tbody>' + "".join(bench_rows) + '</tbody>')

DATA = dict(stories=STORIES, stamp=STAMP, hsbase=HS_BASE)
b64 = base64.b64encode(open(LOGO, "rb").read()).decode()

# ------------------------------------------------------------------ html ----
HTML = r"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Align HCM · Top 30 Command Deck</title>
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&family=Syne:wght@700;800&display=swap');
:root{
  --bg:#060F1D; --bg2:#0A1628; --panel:#0D1B30; --panel-hi:rgba(255,255,255,.035);
  --line:rgba(159,177,199,.13); --line2:rgba(159,177,199,.24);
  --ink:#EAF0F8; --ink2:#9FB1C7; --ink3:#5F7189;
  --orange:#F05A28; --orange-b:#FF6B2B; --steel:#4E86C8; --teal:#1FA98C; --gold:#BA841A; --crit:#E5484D;
  --grad:linear-gradient(135deg,#F05A28 0%,#FF6B35 100%);
  --r:20px;
  --shadow:0 1px 0 rgba(255,255,255,.04) inset, 0 18px 44px rgba(2,8,18,.5);
}
*{box-sizing:border-box;margin:0;padding:0}
::selection{background:rgba(240,90,40,.45);color:#fff}
html{scroll-behavior:smooth}
body{font-family:'Inter',-apple-system,'Segoe UI',Arial,sans-serif;background:var(--bg);color:var(--ink);
  font-size:14px;line-height:1.55;min-height:100vh;-webkit-font-smoothing:antialiased}
body::before{content:"";position:fixed;inset:0;pointer-events:none;z-index:0;opacity:.5;
  background:radial-gradient(820px 560px at 88% -12%, rgba(240,90,40,.16), transparent 60%),
    radial-gradient(700px 520px at -8% 108%, rgba(27,79,114,.30), transparent 62%),
    radial-gradient(1200px 800px at 50% 50%, transparent 55%, rgba(2,6,14,.5) 100%)}
body::after{content:"";position:fixed;inset:0;pointer-events:none;z-index:0;opacity:.045;
  background:linear-gradient(rgba(159,177,199,.7) 1px, transparent 1px) 0 0/100% 44px,
             linear-gradient(90deg, rgba(159,177,199,.7) 1px, transparent 1px) 0 0/44px 100%;
  mask-image:radial-gradient(1100px 700px at 50% 12%, #000 20%, transparent 78%)}
.wrap{max-width:1360px;margin:0 auto;padding:0 28px;position:relative;z-index:1}
a{color:var(--orange-b)}
::-webkit-scrollbar{width:11px;height:11px}
::-webkit-scrollbar-thumb{background:#22344E;border-radius:8px;border:3px solid var(--bg)}
::-webkit-scrollbar-track{background:transparent}
:focus-visible{outline:2px solid var(--orange);outline-offset:2px;border-radius:6px}
.kick{font-size:10px;font-weight:700;letter-spacing:.24em;text-transform:uppercase;color:var(--orange-b)}
h1{font-family:'Syne',sans-serif;font-weight:800;font-size:clamp(25px,3.8vw,44px);letter-spacing:-.015em;line-height:1.05;margin:6px 0 10px;text-wrap:balance}
h1 em{font-style:normal;position:relative;white-space:nowrap}
h1 em::after{content:"";position:absolute;left:0;right:0;bottom:2px;height:5px;border-radius:3px;background:var(--grad);
  transform-origin:left;animation:uline .7s .25s both}
@keyframes uline{from{transform:scaleX(0)}to{transform:scaleX(1)}}
h2{font-family:'Syne',sans-serif;font-weight:800;font-size:clamp(15px,1.6vw,20px);letter-spacing:-.01em}
.sub{color:var(--ink2);font-size:13.5px;max-width:880px;text-wrap:pretty}
.num{font-variant-numeric:tabular-nums}
/* top bar */
header{position:sticky;top:0;z-index:50;background:rgba(6,15,29,.8);backdrop-filter:blur(24px) saturate(1.2);-webkit-backdrop-filter:blur(24px) saturate(1.2);border-bottom:1px solid var(--line)}
header::after{content:"";display:block;height:2px;background:linear-gradient(90deg,transparent,rgba(240,90,40,.7) 30%,rgba(255,107,53,.7) 55%,transparent 85%)}
.bar{display:flex;align-items:center;gap:18px;padding:12px 0;flex-wrap:wrap}
.bar img.logo{width:186px;height:auto;display:block;filter:drop-shadow(0 2px 12px rgba(240,90,40,.25))}
nav{display:flex;gap:2px;flex-wrap:nowrap;background:rgba(255,255,255,.04);border:1px solid var(--line);border-radius:999px;padding:3px;max-width:100%;overflow-x:auto;scrollbar-width:none}
nav::-webkit-scrollbar{display:none}
nav a{font:inherit;font-size:12px;font-weight:600;letter-spacing:.02em;color:var(--ink2);text-decoration:none;white-space:nowrap;
  border-radius:999px;padding:8px 14px;cursor:pointer;transition:color .15s,background .15s}
nav a:hover{color:var(--ink)}
nav a.on{color:#fff;background:var(--grad);box-shadow:0 4px 16px rgba(240,90,40,.4);font-weight:700}
.bar .sp{flex:1}
.stampbox{text-align:right;font-size:9.5px;letter-spacing:.16em;text-transform:uppercase;color:var(--ink3);line-height:1.8}
.livechip{display:inline-flex;align-items:center;gap:6px;font-size:9.5px;font-weight:700;letter-spacing:.16em;color:#33C6A6}
.livechip::before{content:"";width:6px;height:6px;border-radius:50%;background:#33C6A6;box-shadow:0 0 9px #33C6A6;animation:pulse 2.2s infinite}
@keyframes pulse{50%{opacity:.3}}
#q{font:inherit;font-size:13px;color:var(--ink);background:rgba(255,255,255,.045);border:1px solid var(--line2);
  border-radius:999px;padding:9px 16px;min-width:170px}
#q::placeholder{color:var(--ink3)}
#q:focus{outline:none;border-color:var(--orange);box-shadow:0 0 0 3px rgba(240,90,40,.18)}
/* views: stacked without JS, tabbed with JS */
.view{padding:34px 0 30px;scroll-margin-top:86px}
.js .view{display:none;padding-bottom:64px}
.js .view.on{display:block;animation:viewin .35s ease both}
@keyframes viewin{from{opacity:0;transform:translateY(10px)}to{opacity:1;transform:none}}
.shead{display:flex;align-items:baseline;gap:14px;margin:30px 0 14px;flex-wrap:wrap}
.shead small{color:var(--ink3);font-size:12px}
.panel{position:relative;background:linear-gradient(180deg,var(--panel-hi),rgba(255,255,255,0) 55%),var(--panel);
  border-radius:var(--r);padding:20px 22px;box-shadow:var(--shadow)}
.panel::before{content:"";position:absolute;inset:0;border-radius:var(--r);padding:1px;pointer-events:none;
  background:linear-gradient(165deg,rgba(159,177,199,.28),rgba(159,177,199,.08) 38%,rgba(240,90,40,.14) 100%);
  -webkit-mask:linear-gradient(#000 0 0) content-box,linear-gradient(#000 0 0);-webkit-mask-composite:xor;mask-composite:exclude}
.panel>.kick{display:block;margin-bottom:6px}
/* motion guards */
.reveal{opacity:1;transform:none}
.js .reveal{opacity:0;transform:translateY(26px)}
.js .reveal.visible{opacity:1;transform:none;transition:opacity .6s ease,transform .6s cubic-bezier(.2,.7,.2,1)}
.js .d1.visible{transition-delay:.06s}.js .d2.visible{transition-delay:.12s}.js .d3.visible{transition-delay:.18s}
.js .d4.visible{transition-delay:.24s}.js .d5.visible{transition-delay:.30s}
@media (prefers-reduced-motion:reduce){
  html{scroll-behavior:auto}
  .js .reveal,.js .reveal.visible{opacity:1!important;transform:none!important;transition:none!important}
  *,*::before,*::after{animation-duration:.01ms!important;transition-duration:.01ms!important}
}
/* KPI tiles */
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(168px,1fr));gap:12px;margin-top:20px}
.kpi{position:relative;background:linear-gradient(180deg,rgba(255,255,255,.05),rgba(255,255,255,.012)),var(--panel);
  border-radius:16px;padding:15px 17px 13px;box-shadow:var(--shadow);overflow:hidden}
.kpi::before{content:"";position:absolute;inset:0;border-radius:16px;padding:1px;pointer-events:none;
  background:linear-gradient(180deg,rgba(159,177,199,.3),rgba(159,177,199,.07));
  -webkit-mask:linear-gradient(#000 0 0) content-box,linear-gradient(#000 0 0);-webkit-mask-composite:xor;mask-composite:exclude}
.kpi::after{content:"";position:absolute;left:17px;top:0;width:34px;height:3px;border-radius:0 0 3px 3px;background:var(--ac,var(--orange))}
.kpi .lab{font-size:9.5px;font-weight:700;letter-spacing:.18em;text-transform:uppercase;color:var(--ink3);display:block;margin-bottom:7px}
.kpi b{display:block;font-family:'Syne',sans-serif;font-size:25px;font-weight:800;color:#fff;line-height:1.05}
.kpi .ctx{font-size:11px;color:var(--ink2);margin-top:5px}
.kpi.hot b{background:var(--grad);-webkit-background-clip:text;background-clip:text;color:transparent}
/* deck / constellation */
.deck{display:grid;grid-template-columns:minmax(420px,1.12fr) 1fr;gap:16px;align-items:stretch;margin-top:16px}
.orbitbox{position:relative;overflow:hidden;display:flex;align-items:flex-end;justify-content:center;min-height:700px;padding-bottom:26px}
.orbithead{position:absolute;top:18px;left:22px;right:22px;display:flex;justify-content:space-between;align-items:flex-start;gap:10px;z-index:6;flex-wrap:wrap}
.orbitlegend{display:flex;flex-direction:column;gap:6px;align-items:flex-end}
.olg{display:inline-flex;align-items:center;gap:7px;font-size:10.5px;font-weight:600;color:var(--ink2);
  background:rgba(6,15,29,.72);border:1px solid var(--line);border-radius:999px;padding:4px 11px;backdrop-filter:blur(8px)}
.olg .dot{width:8px;height:8px;border-radius:50%}
.orbitscale{width:600px;height:600px;flex:0 0 auto;position:relative}
.orbit{position:absolute;inset:0}
.ringpath{position:absolute;border:1.5px dashed rgba(159,177,199,.16);border-radius:50%;animation:ringspin var(--rd,240s) linear infinite}
@keyframes ringspin{to{transform:rotate(360deg)}}
.coreglow{position:absolute;top:50%;left:50%;width:340px;height:340px;transform:translate(-50%,-50%);border-radius:50%;
  background:radial-gradient(circle,rgba(240,90,40,.22),rgba(240,90,40,.05) 48%,transparent 70%);filter:blur(6px)}
.ring{position:absolute;inset:0;animation:spin var(--dur) linear infinite}
.ring .node{animation:counterspin var(--dur) linear infinite}
@keyframes spin{to{transform:rotate(360deg)}}
@keyframes counterspin{from{transform:translate(-50%,-50%) rotate(0)}to{transform:translate(-50%,-50%) rotate(-360deg)}}
.orbit:hover .ring,.orbit:hover .ring .node,.orbit:hover .ringpath{animation-play-state:paused}
@media (prefers-reduced-motion:reduce){.ring,.ring .node,.ringpath{animation:none}}
.node{position:absolute;top:50%;left:50%;transform:translate(-50%,-50%);width:var(--s);height:var(--s);
  border-radius:13px;background:#fff;border:2px solid var(--hc,rgba(255,255,255,.7));display:flex;align-items:center;justify-content:center;
  cursor:pointer;box-shadow:0 0 0 4px rgba(6,15,29,.55), 0 6px 18px rgba(2,8,18,.5), 0 0 18px -4px var(--hc)}
.node img{max-width:76%;max-height:70%;object-fit:contain}
.node i{display:none;font-style:normal;font-weight:800;font-family:'Syne',sans-serif;color:#0A1628;font-size:calc(var(--s)*.32)}
.node.fb img{display:none}.node.fb i{display:block}
.node:hover{box-shadow:0 0 0 4px rgba(6,15,29,.55), 0 0 0 7px var(--hc), 0 10px 26px rgba(2,8,18,.6);z-index:5}
.core{position:absolute;top:50%;left:50%;transform:translate(-50%,-50%);text-align:center;z-index:2;pointer-events:none}
.core img{width:136px;height:auto;filter:drop-shadow(0 0 26px rgba(240,90,40,.5))}
.core b{display:block;font-family:'Syne',sans-serif;font-size:31px;font-weight:800;margin-top:8px;letter-spacing:-.01em}
.core span{display:block;font-size:9px;letter-spacing:.26em;color:var(--ink2);text-transform:uppercase;margin-top:3px}
.rightcol{display:flex;flex-direction:column;gap:14px}
/* signals */
.feed{display:flex;flex-direction:column;gap:10px;margin-top:12px}
.sig{display:grid;grid-template-columns:10px 1fr auto;gap:12px;align-items:start;background:rgba(255,255,255,.028);
  border:1px solid var(--line);border-radius:14px;padding:12px 14px;cursor:pointer;transition:border-color .15s,background .15s}
.sig:hover{border-color:var(--line2);background:rgba(255,255,255,.05)}
.sig .dot{width:8px;height:8px;border-radius:50%;margin-top:5px;background:var(--sc,var(--orange));box-shadow:0 0 10px var(--sc,var(--orange))}
.sig .k{font-size:9px;font-weight:700;letter-spacing:.2em;text-transform:uppercase;color:var(--sc,var(--orange-b));display:block;margin-bottom:2px}
.sig p{font-size:12.5px;color:var(--ink2)}
.sig p b{color:var(--ink);font-variant-numeric:tabular-nums}
.sig .go{font-size:16px;color:var(--ink3);align-self:center}
/* charts */
.chart{margin-top:10px}
.brow{display:grid;grid-template-columns:172px 1fr 88px;align-items:center;gap:12px;padding:5px 0}
.brow .lab{font-size:12px;font-weight:500;color:var(--ink2);white-space:nowrap;overflow:hidden;text-overflow:ellipsis;text-align:right}
.brow .val{font-size:12px;font-weight:700;font-variant-numeric:tabular-nums;color:var(--ink)}
.track{height:16px;border-radius:5px;background:rgba(255,255,255,.045);position:relative}
.fill{position:absolute;top:2px;bottom:2px;left:0;border-radius:4px;min-width:2px;transform-origin:left}
.js .reveal .fill{transform:scaleX(0)}
.js .reveal.visible .fill{transform:scaleX(1);transition:transform .8s cubic-bezier(.2,.7,.2,1) .15s}
.fill.o{background:linear-gradient(90deg,#C24E20,#F05A28)} .fill.s{background:linear-gradient(90deg,#3D6DA6,#4E86C8)}
.fill.t{background:linear-gradient(90deg,#178069,#1FA98C)}
.duo .track{height:24px}
.duo .fill.a{top:3px;height:8px}
.duo .fill.b{top:13px;height:8px}
.legendrow{display:flex;gap:16px;flex-wrap:wrap;margin-top:12px;font-size:11px;color:var(--ink2)}
.legendrow .sw{display:inline-block;width:10px;height:10px;border-radius:3px;margin-right:6px;vertical-align:-1px}
.heatwrap{overflow-x:auto}
.heat{display:grid;gap:3px;margin-top:8px;min-width:560px}
.heatrow{display:grid;grid-template-columns:150px repeat(8,1fr) 74px;gap:3px;align-items:center}
.heatrow .hl{font-size:11.5px;font-weight:500;color:var(--ink2);white-space:nowrap;overflow:hidden;text-overflow:ellipsis;text-align:right;padding-right:8px}
.cell{height:21px;border-radius:5px;background:rgba(255,255,255,.04)}
.heathead{font-size:9.5px;font-weight:700;letter-spacing:.06em;color:var(--ink3);text-align:center}
.heattot{font-size:11px;font-weight:700;color:var(--ink2);text-align:right;font-variant-numeric:tabular-nums}
#tip{position:fixed;z-index:99;background:rgba(5,12,24,.96);border:1px solid var(--line2);border-radius:12px;padding:10px 13px;
  font-size:12px;color:var(--ink);pointer-events:none;display:none;max-width:300px;box-shadow:0 12px 34px rgba(0,0,0,.55);backdrop-filter:blur(10px)}
#tip b{color:#fff}#tip .m{color:var(--ink2)}
/* deals board */
.dealrow{display:grid;grid-template-columns:44px 1fr auto;gap:14px;align-items:center;padding:12px 4px;border-bottom:1px solid var(--line)}
.dealrow:last-of-type{border-bottom:0}
.dealrow .dn{font-size:13px;font-weight:600;color:var(--ink)}
.dealrow .dm{font-size:11px;color:var(--ink3);margin-top:3px;display:flex;align-items:center;gap:7px;flex-wrap:wrap}
.dealrow .da{text-align:right}
.dealrow .da b{font-family:'Syne',sans-serif;font-size:17px;font-weight:800;color:var(--orange-b)}
.dealrow .da span{display:block;font-size:10px;color:var(--ink3);margin-top:1px}
.dealsum{margin-top:12px;padding-top:12px;border-top:1px solid var(--line);display:flex;justify-content:space-between;align-items:baseline;font-size:12px;color:var(--ink2)}
.dealsum b{font-family:'Syne',sans-serif;font-size:16px;color:var(--orange-b)}
.stage{display:inline-block;font-size:10px;font-weight:700;letter-spacing:.06em;padding:2px 9px;border-radius:999px}
.st-int{background:rgba(78,134,200,.16);color:#9CC0EA}.st-qual{background:rgba(186,132,26,.2);color:#E8B95C}
.st-con{background:rgba(31,169,140,.18);color:#5BD6BC}
/* stories */
.controls{display:flex;gap:8px;align-items:center;margin:16px 0 20px;overflow-x:auto;padding-bottom:4px;scrollbar-width:none}
.controls::-webkit-scrollbar{display:none}
.fbtn{font:inherit;font-size:12px;font-weight:600;padding:8px 15px;border-radius:999px;border:1px solid var(--line2);
  background:rgba(255,255,255,.03);color:var(--ink2);cursor:pointer;transition:all .15s;white-space:nowrap;flex:0 0 auto}
.fbtn:hover{color:var(--ink);border-color:var(--ink3)}
.fbtn.on{background:#fff;color:#0A1628;border-color:#fff;font-weight:700}
.fbtn.orange.on{background:var(--grad);color:#fff;border-color:transparent;box-shadow:0 4px 14px rgba(240,90,40,.35)}
select.fbtn{appearance:none}
.js-only{display:none}
.js .js-only{display:flex}
.grid{display:grid;grid-template-columns:repeat(auto-fill,minmax(min(100%,312px),1fr));gap:14px}
.card{position:relative;background:linear-gradient(180deg,rgba(255,255,255,.035),rgba(255,255,255,0) 60%),var(--panel);
  border-radius:18px;padding:17px;cursor:pointer;transition:transform .16s;box-shadow:var(--shadow)}
.card::before{content:"";position:absolute;inset:0;border-radius:18px;padding:1px;pointer-events:none;
  background:linear-gradient(165deg,rgba(159,177,199,.26),rgba(159,177,199,.07) 45%,rgba(240,90,40,.1));
  -webkit-mask:linear-gradient(#000 0 0) content-box,linear-gradient(#000 0 0);-webkit-mask-composite:xor;mask-composite:exclude}
.card::after{content:"";position:absolute;inset:0;border-radius:18px;padding:1.5px;pointer-events:none;opacity:0;transition:opacity .16s;
  background:linear-gradient(165deg,rgba(255,107,53,.85),rgba(240,90,40,.25) 55%,rgba(255,107,53,.5));
  -webkit-mask:linear-gradient(#000 0 0) content-box,linear-gradient(#000 0 0);-webkit-mask-composite:xor;mask-composite:exclude}
.card:hover{transform:translateY(-3px)}
.card:hover::after{opacity:1}
.card .top{display:flex;gap:12px;align-items:center;padding-right:50px}
.lg{--sz:46px;width:var(--sz);height:var(--sz);flex:0 0 var(--sz);border-radius:12px;background:#fff;
  display:flex;align-items:center;justify-content:center;overflow:hidden;box-shadow:0 0 0 1px rgba(255,255,255,.28), 0 4px 12px rgba(2,8,18,.4)}
.lg img{max-width:78%;max-height:72%;object-fit:contain}
.lg i{display:none;font-style:normal;font-weight:800;font-family:'Syne',sans-serif;color:#0A1628;font-size:calc(var(--sz)*.34)}
.lg.fb img{display:none}.lg.fb i{display:block}
.card h3{font-size:15px;color:#fff;line-height:1.2;font-weight:700}
.card .sid{font-size:10.5px;color:var(--ink3);letter-spacing:.06em;margin-top:2px;word-break:break-all}
.card .rank{position:absolute;top:13px;right:13px;font-family:'Syne',sans-serif;font-weight:800;font-size:12.5px;
  color:#9CC0EA;background:rgba(78,134,200,.13);border:1px solid rgba(78,134,200,.25);border-radius:999px;padding:3px 10px}
.card .rank.top{color:#fff;background:var(--grad);border-color:transparent;box-shadow:0 3px 10px rgba(240,90,40,.4)}
.card .money{font-family:'Syne',sans-serif;font-weight:800;font-size:24px;margin:11px 0 1px;letter-spacing:-.01em}
.card .vbar{height:3px;border-radius:2px;background:rgba(255,255,255,.06);margin:6px 0 9px;overflow:hidden}
.card .vbar i{display:block;height:100%;background:var(--grad);border-radius:2px}
.mline{font-size:11.5px;color:var(--ink2);margin-bottom:10px}
.chips{display:flex;gap:6px;flex-wrap:wrap}
.chip{font-size:10.5px;font-weight:600;padding:3px 9px;border-radius:999px;white-space:nowrap;border:1px solid transparent}
.c-pro{background:rgba(78,134,200,.13);color:#9CC0EA;border-color:rgba(78,134,200,.3)}
.c-ready{background:rgba(31,169,140,.13);color:#5BD6BC;border-color:rgba(31,169,140,.3)}
.c-day{background:rgba(240,90,40,.13);color:#FF9E6E;border-color:rgba(240,90,40,.3)}
.c-ok{background:rgba(31,169,140,.13);color:#5BD6BC;border-color:rgba(31,169,140,.3)}
.c-todo{background:rgba(240,90,40,.12);color:#FF9E6E;border-color:rgba(240,90,40,.28)}
.c-fresh{background:rgba(78,134,200,.13);color:#9CC0EA;border-color:rgba(78,134,200,.3)}
.c-warn{background:rgba(186,132,26,.16);color:#E8B95C;border-color:rgba(186,132,26,.35)}
.c-crit{background:rgba(229,72,77,.13);color:#F58A8E;border-color:rgba(229,72,77,.32)}
.c-mut{background:rgba(255,255,255,.05);color:var(--ink2);border-color:var(--line)}
.hsrow{margin-top:12px;padding-top:11px;border-top:1px dashed var(--line);display:flex;gap:14px;font-size:11px;color:var(--ink2);flex-wrap:wrap;align-items:center}
.hsrow b{color:var(--ink);font-variant-numeric:tabular-nums;font-weight:700}
.hsrow .open b{color:var(--orange-b)}
/* drawer */
#ov{position:fixed;inset:0;background:rgba(2,7,15,.66);backdrop-filter:blur(4px);z-index:80;display:none}
#drawer{position:fixed;top:0;right:-110%;width:min(580px,100vw);height:100dvh;background:linear-gradient(180deg,#0B1A30,#081221);
  border-left:1px solid var(--line2);z-index:90;transition:right .28s cubic-bezier(.2,.7,.2,1);overflow-y:auto;padding:0 24px 32px;overscroll-behavior:contain}
#drawer.on{right:0}
.dband{margin:0 -24px;padding:22px 24px 16px;background:linear-gradient(180deg,rgba(240,90,40,.14),transparent);border-bottom:1px solid var(--line);position:relative}
#drawer .close{position:absolute;top:14px;right:16px;font:inherit;background:rgba(255,255,255,.05);color:var(--ink2);border:1px solid var(--line2);
  border-radius:999px;width:38px;height:38px;cursor:pointer;font-size:15px}
#drawer .close:hover{color:#fff;border-color:var(--ink3)}
.dhead{display:flex;gap:14px;align-items:center;padding-right:44px}
.dgrid{display:grid;grid-template-columns:1fr 1fr;gap:10px;margin:16px 0}
@media (max-width:420px){.dgrid{grid-template-columns:1fr}}
.dcell{background:rgba(255,255,255,.03);border:1px solid var(--line);border-radius:12px;padding:10px 12px}
.dcell span{display:block;font-size:9px;letter-spacing:.16em;text-transform:uppercase;color:var(--ink3);margin-bottom:2px}
.dcell b{font-size:13.5px;color:var(--ink)}
.hslink{display:inline-block;margin:10px 0;font-size:12.5px;font-weight:700;text-decoration:none;color:#fff;
  background:var(--grad);border-radius:999px;padding:11px 20px;box-shadow:0 6px 18px rgba(240,90,40,.35)}
.ministrip{display:flex;gap:4px;margin-top:10px}
.ministrip .yc{flex:1;text-align:center}
.ministrip .yb{height:38px;border-radius:6px;background:rgba(255,255,255,.045);position:relative;overflow:hidden}
.ministrip .yb i{position:absolute;left:0;right:0;bottom:0;background:linear-gradient(180deg,#FF6B35,#C24E20)}
.ministrip span{font-size:9px;color:var(--ink3)}
.checkrow{display:flex;align-items:center;gap:10px;background:rgba(255,255,255,.03);border:1px solid var(--line);border-radius:12px;padding:12px;margin-top:8px;font-size:13px}
.checkrow input{width:19px;height:19px;accent-color:#F05A28}
#drawer textarea{width:100%;min-height:70px;background:rgba(255,255,255,.03);border:1px solid var(--line2);border-radius:12px;
  color:var(--ink);font:inherit;font-size:12.5px;padding:10px;margin-top:8px}
#drawer h2.dsec{font-size:11px;letter-spacing:.2em;text-transform:uppercase;color:var(--orange-b);font-family:'Inter',sans-serif;font-weight:700;margin-top:20px}
/* ops */
.alert{display:grid;grid-template-columns:auto 1fr;gap:14px;align-items:start;background:linear-gradient(90deg,rgba(229,72,77,.15),rgba(229,72,77,.03));
  border:1px solid rgba(229,72,77,.38);border-radius:16px;padding:15px 18px;margin:18px 0}
.alert::before{content:"";width:8px;height:8px;border-radius:50%;background:var(--crit);box-shadow:0 0 12px var(--crit);margin-top:6px}
.alert b{color:#F58A8E}
.alert p{font-size:12.5px;color:var(--ink2);margin-top:3px}
.owners{display:grid;grid-template-columns:repeat(auto-fill,minmax(min(100%,330px),1fr));gap:14px}
.ocard .oh{display:flex;align-items:center;gap:12px;margin-bottom:12px}
.avatar{width:42px;height:42px;border-radius:50%;background:var(--grad);display:flex;align-items:center;justify-content:center;
  font-family:'Syne',sans-serif;font-weight:800;color:#fff;font-size:14px;box-shadow:0 4px 12px rgba(240,90,40,.3);flex:0 0 42px}
.avatar.gray{background:linear-gradient(135deg,#39445A,#22304a);box-shadow:none}
.ocard h3{font-size:14.5px;font-weight:700}
.ocard .ost{font-size:10.5px;color:var(--ink3);margin-top:1px}
.olist{display:flex;flex-direction:column;gap:6px}
.oitem{display:flex;align-items:center;gap:8px;font-size:12px;color:var(--ink2);background:rgba(255,255,255,.028);
  border:1px solid var(--line);border-radius:10px;padding:8px 11px;flex-wrap:wrap}
.oitem b{color:var(--ink);font-weight:700}
.oitem .sp{flex:1}
.copybtn{font:inherit;font-size:11px;font-weight:700;color:var(--ink);background:rgba(255,255,255,.04);border:1px solid var(--line2);
  border-radius:999px;padding:9px 16px;cursor:pointer;margin-top:12px;transition:all .15s}
.copybtn:hover{border-color:var(--orange);color:var(--orange-b)}
.meter{height:10px;border-radius:6px;background:rgba(255,255,255,.06);overflow:hidden;margin-top:10px}
.meter i{display:block;height:100%;background:var(--grad);border-radius:6px;transition:width .5s cubic-bezier(.2,.7,.2,1)}
/* coverage / bench / foot */
.covgrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(min(100%,460px),1fr));gap:14px}
.gaprow{display:grid;grid-template-columns:148px 1fr 58px 86px;gap:10px;align-items:center;padding:4.5px 0;font-size:12px}
.gaprow .lab{color:var(--ink2);text-align:right;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;font-weight:500}
.gaprow .n{color:var(--ink3);font-variant-numeric:tabular-nums}
.benchwrap{overflow-x:auto;margin-top:16px}
table.bench{width:100%;border-collapse:collapse;min-width:760px}
table.bench th{font-size:9.5px;letter-spacing:.14em;text-transform:uppercase;color:var(--ink3);text-align:left;padding:10px;border-bottom:1px solid var(--line2)}
table.bench td{padding:10px;border-bottom:1px solid var(--line);font-size:12.5px;color:var(--ink2)}
table.bench td b{color:var(--ink)}
table.bench tr:hover td{background:rgba(255,255,255,.025)}
table.bench tr:last-child td{border-bottom:0}
.tier{font-size:10.5px;font-weight:700;padding:3px 10px;border-radius:999px;white-space:nowrap}
.tier.a{background:rgba(31,169,140,.14);color:#5BD6BC}.tier.b{background:rgba(78,134,200,.14);color:#9CC0EA}
.foot{border-top:1px solid var(--line);margin-top:36px;padding:28px 0 44px;display:flex;gap:18px;align-items:center;justify-content:space-between;flex-wrap:wrap}
.foot img{width:150px;filter:drop-shadow(0 2px 10px rgba(240,90,40,.2))}
.foot p{font-size:10.5px;color:var(--ink3);text-align:right;line-height:1.8}
/* ---------------- mobile ---------------- */
@media (max-width:1060px){.deck{grid-template-columns:1fr}}
@media (max-width:1240px){
  .orbitscale{transform:scale(.85);transform-origin:bottom center}
  .orbitbox{min-height:620px}
}
@media (max-width:900px){.stampbox{display:none}}
@media (max-width:760px){
  .wrap{padding:0 16px}
  .bar{gap:10px;padding:10px 0}
  .bar img.logo{width:142px}
  nav{order:3;width:100%}
  #q{min-width:0;flex:1;font-size:16px}
  .view{padding:22px 0 24px;scroll-margin-top:118px}
  .js .view.on{padding-bottom:48px}
  .kpis{grid-template-columns:repeat(2,1fr);gap:9px}
  .kpi{padding:12px 13px 11px}
  .kpi b{font-size:20px}
  .orbitscale{transform:scale(.56);transform-origin:bottom center}
  .orbitbox{min-height:430px;padding-bottom:10px}
  .orbithead{position:static;margin:16px 16px 0;order:-1}
  .orbitbox{flex-direction:column;align-items:center}
  .orbitlegend{flex-direction:row;flex-wrap:wrap;align-items:flex-start;justify-content:flex-start;gap:5px;margin-top:8px}
  .core img{width:110px}
  .core b{font-size:24px}
  .brow{grid-template-columns:96px 1fr 64px;gap:8px}
  .brow .lab{font-size:10.5px}
  .brow .val{font-size:10.5px}
  .heatrow{grid-template-columns:86px repeat(8,1fr) 56px}
  .heatrow .hl{font-size:10px}
  .cell{height:16px}
  .heattot{font-size:9.5px}
  .gaprow{grid-template-columns:96px 1fr 44px 70px;gap:6px;font-size:11px}
  .gaprow .chip{font-size:9px;padding:2px 7px}
  .dealrow{grid-template-columns:36px 1fr;gap:10px}
  .dealrow .lg{--sz:36px}
  .dealrow .da{grid-column:2;text-align:left;display:flex;gap:10px;align-items:baseline;margin-top:2px}
  .dealrow .da b{font-size:15px}
  .panel{padding:16px}
  .foot{flex-direction:column;align-items:flex-start}
  .foot p{text-align:left}
  h1{font-size:25px}
}
@media (max-width:420px){
  .orbitscale{transform:scale(.5)}
  .orbitbox{min-height:395px}
  .kpis{grid-template-columns:1fr 1fr}
}
@media print{header nav,#q,.controls,.copybtn{display:none}.view{display:block!important;padding:12px 0}body::before,body::after{display:none}}
</style></head><body data-palette="#F05A28,#4E86C8,#1FA98C,#BA841A">
<script>document.documentElement.classList.add('js')</script>
<div id="tip"></div>
<header><div class="wrap bar">
  <img class="logo" src="data:image/png;base64,@@LOGO@@" alt="Align HCM. Human Capital Management.">
  <nav id="nav">
    <a href="#v-overview" data-v="overview" class="on">Command Deck</a>
    <a href="#v-stories" data-v="stories">Stories</a>
    <a href="#v-pipeline" data-v="pipeline">Pipeline · Live</a>
    <a href="#v-ops" data-v="ops">Validation Ops</a>
    <a href="#v-coverage" data-v="coverage">Coverage</a>
    <a href="#v-bench" data-v="bench">Bench</a>
  </nav>
  <div class="sp"></div>
  <input id="q" type="search" placeholder="Search the book…" aria-label="Search">
  <div class="stampbox"><span class="livechip">HubSpot @@PORTAL@@</span><br>Read @@STAMP@@ · read-only</div>
</div></header>

<div class="wrap">

<section class="view on" id="v-overview">
  <span class="kick">Internal · Unpublished · Objective 1</span>
  <h1>Top 30 <em>Command Deck</em></h1>
  <p class="sub">One deck for the story shortlist and the live book behind it. Story data is the 10 Aug build. Every deal figure was read from HubSpot portal @@PORTAL@@ on @@STAMP@@, read-only, with won, open, and lost split by real deal state. Client marks load from each company's public domain.</p>
  <div class="kpis">@@KPIS@@</div>
  <div class="deck">
    <div class="panel orbitbox reveal">
      <div class="orbithead">
        <div><span class="kick">The book</span><h2>Story constellation</h2></div>
        <div class="orbitlegend">@@ORBITLEGEND@@</div>
      </div>
      <div class="orbitscale"><div class="orbit" id="orbit">@@ORBIT@@</div></div>
    </div>
    <div class="rightcol">
      <div class="panel reveal d1">
        <span class="kick">Portfolio signal</span>
        <h2>Live signals</h2>
        <div class="feed" id="feed">@@FEED@@</div>
      </div>
      <div class="panel reveal d2">
        <span class="kick">Magnitude</span>
        <h2>Highest-value stories</h2>
        <div class="chart">@@TOPBARS@@</div>
      </div>
      <div class="panel reveal d3">
        <span class="kick">Readiness</span>
        <h2>Path to public use</h2>
        <div class="chart">@@READYBARS@@</div>
      </div>
    </div>
  </div>
</section>

<section class="view" id="v-stories">
  <span class="kick">The thirty</span>
  <h1>Story <em>cards</em></h1>
  <p class="sub">Tap any card for the full record: story status, CRM footprint with win rate, closed-won history, open deals, validation checklist. Filters stack with the global search.</p>
  <div class="controls js-only" id="storyfilters">
    <button class="fbtn orange on" data-f="plat" data-val="">All platforms</button>
    <button class="fbtn orange" data-f="plat" data-val="UKG Pro">UKG Pro</button>
    <button class="fbtn orange" data-f="plat" data-val="UKG Ready">UKG Ready</button>
    <button class="fbtn orange" data-f="plat" data-val="Dayforce">Dayforce</button>
    <button class="fbtn on" data-f="ready" data-val="">Any status</button>
    <button class="fbtn" data-f="ready" data-val="ready">Ready</button>
    <button class="fbtn" data-f="ready" data-val="fresh">Fresh quote</button>
    <button class="fbtn" data-f="ready" data-val="todo">Needs narrative</button>
    <select class="fbtn" id="sortsel">
      <option value="rank">Sort · Rank</option>
      <option value="value">Sort · Story value</option>
      <option value="open">Sort · Open pipeline</option>
      <option value="won">Sort · Lifetime won</option>
      <option value="client">Sort · Client A to Z</option>
    </select>
    <button class="fbtn" id="csv">⬇ CSV</button>
  </div>
  <div class="grid" id="cards">@@CARDS@@</div>
</section>

<section class="view" id="v-pipeline">
  <span class="kick">CRM truth · read @@STAMP@@</span>
  <h1>Live <em>pipeline</em></h1>
  <p class="sub">28 of 30 clients matched to HubSpot company records by domain. JACAM Catalyst and Ace Parking carry story data only; their CRM records have no domain. Won, open, and lost are split by real deal state: open means not yet closed, and lost is closed business that did not land.</p>
  <div class="kpis">@@PULSE@@</div>
  <div class="shead"><span class="kick">Money in motion</span><h2>Open deals on shortlist clients</h2><small>Every open deal, live stage and owner · sorted by size</small></div>
  <div class="panel reveal">@@DEALSBOARD@@</div>
  <div class="shead"><span class="kick">The relationship ledger</span><h2>Won vs open, per client</h2><small>Won in steel, open in orange · lost shown on hover</small></div>
  <div class="panel reveal">
    <div class="legendrow"><span><span class="sw" style="background:var(--steel)"></span>Closed won, lifetime</span><span><span class="sw" style="background:var(--orange)"></span>Open now</span></div>
    <div class="chart duo">@@WONOPEN@@</div>
  </div>
  <div class="shead"><span class="kick">Eight years of delivery</span><h2>Closed-won history, 2019 to 2026</h2><small>Deals won per client per year · shade is won value</small></div>
  <div class="panel reveal"><div class="heatwrap"><div class="heat">@@HEAT@@</div></div></div>
</section>

<section class="view" id="v-ops">
  <span class="kick">Objective 3 · clearance</span>
  <h1>Validation <em>ops</em></h1>
  <p class="sub">Reference Status, Validator, and Validation Notes are blank on all 453 source rows. This board turns the 30 into owner-sized ask lists. Checklist state saves in this browser only.</p>
  <div class="alert">@@ALERT@@</div>
  <div class="panel reveal" style="margin:0 0 16px">
    <span class="kick">Progress</span>
    <b id="progresslab" style="font-family:Syne;font-size:17px;display:block;margin-top:4px">0 of 30 stories fully validated</b>
    <div class="meter"><i id="progressbar" style="width:0%"></i></div>
  </div>
  <div class="owners" id="owners">@@OWNERS@@</div>
</section>

<section class="view" id="v-coverage">
  <span class="kick">Objective 4 · representation</span>
  <h1>Coverage <em>radar</em></h1>
  <p class="sub">Shortlist coverage against the whole 453-row book, by dimension. Bars are book value; the count shows how many shortlist stories cover that slice. Flags mark unrepresented dimensions.</p>
  <div class="covgrid" style="margin-top:16px">@@COV@@</div>
</section>

<section class="view" id="v-bench">
  <span class="kick">Objective 2 · references</span>
  <h1>Reference <em>bench</em></h1>
  <p class="sub">Ranked by how close each client already is to usable: Tier A is Raven plus recent, Tier B is Raven but older. Sheet value sits beside live CRM lifetime won for cross-checking.</p>
  <div class="panel reveal benchwrap"><table class="bench">@@BENCH@@</table></div>
</section>

<div class="foot">
  <img src="data:image/png;base64,@@LOGO@@" alt="Align HCM">
  <p>Align HCM · SmartCare · UKG implementation and support · Internal use, unpublished<br>
  Story data: 10 Aug 2026 shortlist build · CRM reads: HubSpot portal @@PORTAL@@, @@STAMP@@, read-only · Client marks load from public domains at view time</p>
</div>
</div>

<div id="ov"></div>
<aside id="drawer" aria-label="Story detail"></aside>

<script>
const DATA = @@DATA@@;
const HSB = DATA.hsbase;
const $ = s => document.querySelector(s), $$ = s => [...document.querySelectorAll(s)];
const fmt = n => "$" + Math.round(n).toLocaleString("en-US");
const esc = s => String(s).replace(/[&<>"]/g, c => ({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"}[c]));
const initials = n => n.replace(/&/g," ").split(/\s+/).filter(w=>/[A-Za-z0-9]/.test(w[0])).slice(0,2).map(w=>w[0]).join("").toUpperCase();
const readiness = s => s.needs.startsWith("Ready") ? "ready" : s.needs.startsWith("Fresh") ? "fresh" : "todo";
const READY_LABEL = {ready:"Ready · spot-check", fresh:"Needs fresh quote", todo:"Needs narrative + quote"};
const PLATC = {"UKG Pro":"c-pro","UKG Ready":"c-ready","Dayforce":"c-day"};
const STAGEC = {"Expressing Interest":"st-int","Qualification":"st-qual","Contracting":"st-con"};
const S = DATA.stories;
const logoChip = (s, sz) =>
  `<span class="lg" style="--sz:${sz}px"><img loading="lazy" src="https://logo.clearbit.com/${s.domain}" alt="" onerror="this.parentElement.classList.add('fb')"><i>${initials(s.client)}</i></span>`;

/* tooltip: hover + tap */
const tip = $("#tip");
function showTip(el, x, y){
  tip.innerHTML = el.dataset.tip; tip.style.display = "block";
  const w = tip.offsetWidth, h = tip.offsetHeight;
  tip.style.left = Math.max(6, Math.min(x+14, innerWidth-w-8)) + "px";
  tip.style.top = Math.min(y+16, innerHeight-h-8) + "px";
}
document.addEventListener("mousemove", e=>{
  const el = e.target.closest("[data-tip]");
  if(el) showTip(el, e.clientX, e.clientY); else tip.style.display="none";
});
document.addEventListener("touchstart", e=>{
  const el = e.target.closest("[data-tip]");
  if(el && !e.target.closest(".card") && !e.target.closest(".node")){
    const t = e.touches[0]; showTip(el, t.clientX, t.clientY-40);
    clearTimeout(showTip._t); showTip._t = setTimeout(()=>tip.style.display="none", 2400);
  } else tip.style.display="none";
}, {passive:true});

/* reveal */
const io = new IntersectionObserver(es=>es.forEach(e=>{ if(e.isIntersecting){ e.target.classList.add("visible"); io.unobserve(e.target);} }),{threshold:.1});
function arm(root=document){ root.querySelectorAll(".reveal:not(.visible)").forEach(el=>io.observe(el)); }
/* count-up */
const reduced = matchMedia("(prefers-reduced-motion: reduce)").matches;
function countUp(el){
  if(reduced || el._done) return; el._done = true;
  const fin = el.dataset.final, m = fin.match(/^([^0-9]*)([\d,.]+)(.*)$/);
  if(!m) return;
  const num = parseFloat(m[2].replace(/,/g,"")), dec = (m[2].split(".")[1]||"").length;
  const t0 = performance.now(), D = 850;
  (function tick(t){
    const p = Math.min(1,(t-t0)/D), v = num*(1-Math.pow(1-p,3));
    el.textContent = m[1] + v.toLocaleString("en-US",{minimumFractionDigits:dec,maximumFractionDigits:dec}) + m[3];
    if(p<1) requestAnimationFrame(tick); else el.textContent = fin;
  })(t0);
}
/* tabs */
function show(v){
  $$("nav a").forEach(b=>b.classList.toggle("on", b.dataset.v===v));
  $$(".view").forEach(s=>s.classList.toggle("on", s.id==="v-"+v));
  history.replaceState(null,"","#"+v); window.scrollTo({top:0});
  arm(); $$("#v-"+v+" .kpi b[data-final]").forEach(countUp);
}
$("#nav").addEventListener("click", e=>{
  const a = e.target.closest("a[data-v]"); if(!a) return;
  e.preventDefault(); show(a.dataset.v);
});
addEventListener("keydown", e=>{
  if(e.key==="/" && document.activeElement!==$("#q")){ e.preventDefault(); $("#q").focus(); }
  if(e.key==="Escape") closeDrawer();
});

/* filters, sort, search on server-rendered cards */
const filters = {plat:"", ready:"", q:""};
function applyCards(){
  const cards = $$("#cards .card");
  cards.forEach(c=>{
    const ok = (!filters.plat || c.dataset.plat===filters.plat) &&
      (!filters.ready || c.dataset.ready===filters.ready) &&
      (!filters.q || c.dataset.q.includes(filters.q));
    c.style.display = ok ? "" : "none";
  });
  const m = $("#sortsel").value, box = $("#cards");
  const key = c => m==="rank"? +c.dataset.rank : m==="value"? -c.dataset.value :
    m==="open"? -c.dataset.open : m==="won"? -c.dataset.won : 0;
  [...box.children].sort((a,b)=> m==="client"? a.dataset.client.localeCompare(b.dataset.client) : key(a)-key(b))
    .forEach(c=>box.appendChild(c));
}
$("#storyfilters").addEventListener("click", e=>{
  const b = e.target.closest("[data-f]"); if(!b) return;
  filters[b.dataset.f] = b.dataset.val;
  $$(`#storyfilters [data-f="${b.dataset.f}"]`).forEach(x=>x.classList.toggle("on", x===b));
  applyCards();
});
$("#sortsel").addEventListener("change", applyCards);
$("#q").addEventListener("input", e=>{ filters.q = e.target.value.toLowerCase().trim(); applyCards();
  if(!$("#v-stories").classList.contains("on") && filters.q) show("stories"); });
$("#cards").addEventListener("click", e=>{ const c = e.target.closest(".card"); if(c) openDrawer(c.dataset.client); });
$("#orbit").addEventListener("click", e=>{ const n = e.target.closest(".node"); if(n) openDrawer(n.dataset.client); });
$("#feed").addEventListener("click", e=>{ const g = e.target.closest(".sig"); if(g) show(g.dataset.go); });
$("#csv") && $("#csv").addEventListener("click", ()=>{
  const head = "rank,story_id,client,industry,size_hc,geo,platform,story_type,start_yr,story_value_usd,narrative,quote,needs,crm_owner,owner_active,domain,hq,employees_crm,crm_deals,crm_won_usd,crm_open_usd,crm_lost_usd,win_rate_pct,crm_company_id";
  const rows = S.map(s=>[s.rank,s.sid,`"${s.client}"`,`"${s.industry}"`,`"${s.size}"`,s.geo,`"${s.platform}"`,`"${s.stype}"`,s.yr,s.value,s.narr,s.quote,`"${s.needs}"`,`"${s.owner}"`,s.ownerActive===false?"INACTIVE":"active",s.domain,`"${s.hq||""}"`,s.emp||"",s.hs?s.hs.deals:"",s.hs?s.hs.won:"",s.hs?s.hs.open:"",s.hs?s.hs.lost:"",s.hs&&(s.hs.won_n+s.hs.lost_n)?Math.round(100*s.hs.won_n/(s.hs.won_n+s.hs.lost_n)):"",s.hs?s.hs.crm_id:""].join(","));
  const blob = new Blob([head+"\n"+rows.join("\n")], {type:"text/csv"});
  const a = Object.assign(document.createElement("a"), {href:URL.createObjectURL(blob), download:"align_top30_command_deck.csv"});
  a.click(); URL.revokeObjectURL(a.href);
});

/* drawer */
const store = {
  k: "align-top30-validation",
  read(){ try{ return JSON.parse(localStorage.getItem(this.k))||{} }catch(e){ return {} } },
  write(o){ localStorage.setItem(this.k, JSON.stringify(o)); }
};
function openDrawer(client){
  const s = S.find(x=>x.client===client); if(!s) return;
  const r = readiness(s), st = store.read()[s.sid]||{};
  let strip = "";
  if(s.hs){
    const ys = []; for(let y=2019;y<=2026;y++) ys.push([y, s.hs.years[y]||null]);
    const ymax = Math.max(1,...ys.map(([,v])=>v?v[1]:0));
    strip = `<div class="ministrip">`+ys.map(([y,vv])=>`
      <div class="yc"><div class="yb" data-tip="${vv? `<b>${y}</b><br>${vv[0]} deal${vv[0]>1?"s":""} won · ${fmt(vv[1])}` : `<b>${y}</b><br>No wins`}">
        ${vv? `<i style="height:${Math.max(8,100*vv[1]/ymax)}%"></i>`:""}</div><span>${String(y).slice(2)}</span></div>`).join("")+`</div>`;
  }
  const openList = s.hs && s.hs.deals_open.length ? s.hs.deals_open.map(d=>`
      <div class="dealrow"><span class="lg fb" style="--sz:36px"><i>${initials(s.client)}</i></span>
        <div><div class="dn">${esc(d.name)}</div><div class="dm"><span class="stage ${STAGEC[d.stage]||""}">${esc(d.stage)}</span>${esc(d.owner)} · closes ${d.close}</div></div>
        <div class="da"><b>${fmt(d.amt)}</b></div></div>`).join("") : "";
  const closed = s.hs? s.hs.won_n+s.hs.lost_n : 0;
  $("#drawer").innerHTML = `
    <div class="dband">
      <button class="close" aria-label="Close">✕</button>
      <div class="dhead">${logoChip(s,54)}<div>
        <h2 style="font-size:21px">${esc(s.client)}</h2>
        <div style="color:var(--ink3);font-size:11px;margin-top:2px">#${s.rank} · ${esc(s.sid)} · <a href="https://${s.domain}" target="_blank" rel="noopener">${s.domain}</a></div>
      </div></div>
      <div class="chips" style="margin-top:10px">
        <span class="chip ${PLATC[s.platform]||""}">${esc(s.platform)}</span>
        <span class="chip c-mut">${esc(s.stype)}</span>
        <span class="chip ${r==="ready"?"c-ok":r==="fresh"?"c-fresh":"c-todo"}">${READY_LABEL[r]}</span>
      </div>
    </div>
    <div class="dgrid">
      <div class="dcell"><span>Story value · true USD</span><b style="font-family:Syne;font-size:19px" class="num">${fmt(s.value)}</b></div>
      <div class="dcell"><span>Start year</span><b>${s.yr}</b></div>
      <div class="dcell"><span>Industry</span><b>${esc(s.industry)}</b></div>
      <div class="dcell"><span>Size · Geo</span><b>${esc(s.size)} · ${esc(s.geo)}</b></div>
      <div class="dcell"><span>Narrative</span><b style="color:${s.narr==="Yes"?"#5BD6BC":"#FF9E6E"}">${esc(s.narr)}</b></div>
      <div class="dcell"><span>Quote</span><b style="color:${s.quote==="Raven"?"#5BD6BC":"#FF9E6E"}">${esc(s.quote)}</b></div>
      <div class="dcell"><span>CRM deal owner</span><b>${esc(s.owner)}${s.ownerActive===false?' <span class="chip c-crit">inactive</span>':""}</b></div>
      <div class="dcell"><span>What it still needs</span><b>${esc(s.needs)}</b></div>
      ${s.hq?`<div class="dcell"><span>HQ (CRM)</span><b>${esc(s.hq)}</b></div>`:""}
      ${s.emp?`<div class="dcell"><span>Employees (CRM)</span><b>${Number(s.emp).toLocaleString()}</b></div>`:""}
    </div>
    ${s.hs? `<h2 class="dsec">HubSpot record · live ${DATA.stamp}</h2>
      <div class="dgrid">
        <div class="dcell"><span>CRM company</span><b>${esc(s.hs.crm_name)}</b></div>
        <div class="dcell"><span>Deals on record</span><b>${s.hs.deals}</b></div>
        <div class="dcell"><span>Closed won</span><b style="color:#5BD6BC">${fmt(s.hs.won)} · ${s.hs.won_n}</b></div>
        <div class="dcell"><span>Open now</span><b style="color:#FF9E6E">${s.hs.open? fmt(s.hs.open):"None"}${s.hs.open_n? " · "+s.hs.open_n:""}</b></div>
        <div class="dcell"><span>Closed lost</span><b>${s.hs.lost? fmt(s.hs.lost):"None"}${s.hs.lost_n? " · "+s.hs.lost_n:""}</b></div>
        <div class="dcell"><span>Win rate · closed deals</span><b>${closed? Math.round(100*s.hs.won_n/closed)+"%":"·"}</b></div>
      </div>
      ${openList? `<h2 class="dsec">Open deals</h2>${openList}`:""}
      <h2 class="dsec">Closed-won by year</h2>${strip}
      <a class="hslink" href="${HSB}${s.hs.crm_id}" target="_blank" rel="noopener">Open in HubSpot ↗</a>`
    : `<p style="color:var(--ink3);font-size:12.5px;margin:10px 0">No CRM company record matched by domain. Story data only.</p>`}
    <h2 class="dsec">Validation checklist <span style="font-size:9px;color:var(--ink3);letter-spacing:.1em">· saved in this browser</span></h2>
    ${["Narrative approved","Quote sourced","Reference status set","Client cleared for public naming"].map((lab,i)=>`
      <label class="checkrow"><input type="checkbox" data-ck="${i}" ${st["c"+i]?"checked":""}> ${lab}</label>`).join("")}
    <textarea placeholder="Validation notes for ${esc(s.client)}…" data-note>${esc(st.note||"")}</textarea>`;
  $("#drawer").classList.add("on"); $("#ov").style.display="block";
  $("#drawer .close").onclick = closeDrawer; $("#ov").onclick = closeDrawer;
  $$("#drawer [data-ck]").forEach(cb=>cb.addEventListener("change", ()=>{
    const o = store.read(); o[s.sid] = o[s.sid]||{}; o[s.sid]["c"+cb.dataset.ck] = cb.checked; store.write(o); refreshOps();
  }));
  $("#drawer [data-note]").addEventListener("input", e=>{
    const o = store.read(); o[s.sid] = o[s.sid]||{}; o[s.sid].note = e.target.value; store.write(o);
  });
}
function closeDrawer(){ $("#drawer").classList.remove("on"); $("#ov").style.display="none"; }

/* ops state from localStorage */
function refreshOps(){
  const st = store.read();
  const done = S.filter(s=>{ const o=st[s.sid]; return o&&o.c0&&o.c1&&o.c2&&o.c3; }).length;
  $("#progresslab").textContent = `${done} of 30 stories fully validated`;
  $("#progressbar").style.width = (100*done/30)+"%";
  $$("#owners .oitem").forEach(it=>{
    const o = st[it.dataset.sid], k = o&&o.c0&&o.c1&&o.c2&&o.c3;
    const chip = it.querySelector(".ost-chip");
    if(k){ chip.textContent = "validated"; chip.className = "chip c-ok ost-chip"; }
    else { chip.textContent = it.dataset.needs; }
  });
}
$$("#owners .copybtn").forEach(b=>b.addEventListener("click", ()=>{
  const list = S.filter(s=>s.owner===b.dataset.owner);
  const txt = `Validation asks for ${b.dataset.owner} (Align Top 30, ${DATA.stamp}):\n` +
    list.map(s=>`• #${s.rank} ${s.client} (${s.sid}): ${s.needs}. Story value ${fmt(s.value)}.`).join("\n");
  navigator.clipboard.writeText(txt).then(()=>{ b.textContent="Copied ✓"; setTimeout(()=>b.textContent="Copy ask list",1600); });
}));

refreshOps(); arm();
$$("#v-overview .kpi b[data-final]").forEach(countUp);
if(location.hash){ const v = location.hash.slice(1).replace(/^v-/,""); if($("#v-"+v)) show(v); }
</script>
</body></html>"""

page = (HTML.replace("@@DATA@@", json.dumps(DATA, separators=(",",":")))
    .replace("@@KPIS@@", KPIS).replace("@@ORBITLEGEND@@", ORBITLEGEND).replace("@@ORBIT@@", ORBIT)
    .replace("@@FEED@@", FEED).replace("@@TOPBARS@@", TOPBARS).replace("@@READYBARS@@", READYBARS)
    .replace("@@CARDS@@", CARDS).replace("@@PULSE@@", PULSE_HTML).replace("@@DEALSBOARD@@", DEALSBOARD)
    .replace("@@WONOPEN@@", WONOPEN).replace("@@HEAT@@", HEAT).replace("@@ALERT@@", ALERT)
    .replace("@@OWNERS@@", OWNERS_HTML).replace("@@COV@@", COV).replace("@@BENCH@@", BENCH)
    .replace("@@LOGO@@", b64).replace("@@STAMP@@", STAMP).replace("@@PORTAL@@", PORTAL))
os.makedirs(os.path.dirname(OUT), exist_ok=True)
open(OUT, "w").write(page)
print("wrote", OUT, f"{len(page)/1024:.0f} KB")
print("server-rendered:", all(x not in page for x in ("@@KPIS@@","@@CARDS@@","@@HEAT@@","@@OWNERS@@")))
