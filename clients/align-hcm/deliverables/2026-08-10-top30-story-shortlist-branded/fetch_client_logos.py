from __future__ import annotations

import io
import json
import re
import time
from pathlib import Path
from urllib.parse import urljoin, urlparse

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook
from PIL import Image, ImageChops

try:
    import cairosvg
except Exception:
    cairosvg = None


ROOT = Path(__file__).resolve().parent
SOURCE = next(
    path
    for path in (
        ROOT / "upload" / "Align_HCM_Top30_Story_Shortlist_Branded.xlsx",
        ROOT / "Align_HCM_Top30_Story_Shortlist_Branded.xlsx",
    )
    if path.exists()
)
OUT = ROOT / "tmp" / "pdfs" / "client-logos"
OUT.mkdir(parents=True, exist_ok=True)

SESSION = requests.Session()
SESSION.headers.update(
    {
        "User-Agent": "Mozilla/5.0 (compatible; AlignHCMEditorial/1.0; +https://alignhcm.com)",
        "Accept": "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8",
    }
)

VERIFIED_OVERRIDES = {
    "bayshore.ca": [
        "https://www.bayshore.ca/wp-content/themes/bayshore/dist/images/logo.svg",
    ],
    "cccu.ca": [
        "https://www.cccu.ca/assets/img/logos/6954ba2a-6878-4812-a10d-c52332946014.svg",
    ],
    "chq.org": [
        "https://www.chq.org/wp-content/uploads/2021/05/ChautauquaLogo_2021_PMS348.svg",
    ],
    "everquote.com": [
        "https://cdn.prod.website-files.com/5f19961dcd0dc8de84b6d0c9/5f19b00072b3350a15873d03_EverQuote%20Logo.svg",
    ],
    "mta.info": [
        "https://upload.wikimedia.org/wikipedia/commons/3/3c/MTA_NYC_logo.svg",
    ],
    "mwcomponents.com": [
        "https://mma.prnewswire.com/media/1500972/MW_Components_Logo.jpg",
    ],
    "trimac.com": [
        "https://dokumfe7mps0i.cloudfront.net/oms/2931/image/2025/3/DNV4A_google/google.jpg",
        "https://www.trimac.com/wp-content/uploads/2025/12/Trimac-logo.svg",
    ],
    "tsh.org": [
        "https://portal.tristatehospital.org/Pub/Phm/Images/v3/TRSWEB.LIVEF/SignOnImage.png",
        "https://www.tsh.org/wp-content/uploads/2020/10/logo.svg",
    ],
    "rollins.com": [
        "https://d1io3yog0oux5.cloudfront.net/_d4e6516391e7a29b8f572b54c9bd7951/rollins/db/1002/9855/image.svg",
    ],
    "eaglematerials.com": [
        "https://assets1.hbsdealer.com/images/v/max_width_1440/s3fs-public/2024-05/eagle.png",
        "https://www.eaglematerials.com/sites/default/files/eagle-materials.png",
    ],
    "grandriverhealth.org": [
        "https://grandriverhealth.org/wp-content/uploads/2022/10/GRH-main-logo-2022.png",
    ],
    "mccain.com": [
        "https://www.mccain.com/media/3565/mccain-foods-logo.jpg",
    ],
    "rei.com": [
        "https://www.rei.com/rei-co-op-logo-black.svg",
    ],
    "troon.com": [
        "https://bayclub.onelombard.com/wp-content/uploads/sites/4/2024/01/troon-vector-logo.png",
    ],
    "ges.com": [
        "https://mma.prnewswire.com/media/2426524/GES_COLOR_Logo.jpg",
    ],
    "burnco.com": [
        "https://cdn.craft.cloud/284c7dbf-9389-4754-b7ea-50312329b835/assets/contentful/BURNCO-Pantone-485.jpg?fit=cover&format=webp&s=n961Eg_CNtjmIxeOIicr91ryKBGxXuGnT_HV09WvLp8",
    ],
    "primeflight.com": [
        "https://i.vimeocdn.com/portrait/36472638_640x640?region=us&sig=6cbda6caed1c9f89643da948869d6445173179788e8fd166b5c76a53ad6a0652&v=1",
    ],
    "distributorwireandcable.com": [
        "https://www.distributorwire.com/images/seo/DWC-meta-1a.png",
    ],
    "aceparking.com": [
        "https://www.sandiego.org/sites/default/files/styles/large/public/listing_images/sandiego-d44206053d5d44868d1b5d45ead1097f_5A160198-0725-DE67-D29B7B99CEF51A75-5a1600f6f00bcd1_5a160f51-b81e-21e9-b468ddfc2f78fcf6.jpg.webp?itok=kY7s8HW0",
        "https://icon.horse/icon/aceparking.com",
        "https://icons.duckduckgo.com/ip3/aceparking.com.ico",
    ],
}


def slug(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", "-", value.lower()).strip("-")


def image_from_response(response: requests.Response) -> Image.Image:
    payload = response.content
    content_type = (response.headers.get("content-type") or "").lower()
    url_path = urlparse(response.url).path.lower()
    is_svg = "svg" in content_type or url_path.endswith(".svg") or payload.lstrip().startswith(b"<svg")
    if is_svg:
        if not cairosvg:
            raise ValueError("SVG support unavailable")
        payload = cairosvg.svg2png(bytestring=payload, output_width=800)
    image = Image.open(io.BytesIO(payload)).convert("RGBA")
    if image.width < 16 or image.height < 16:
        raise ValueError("image too small")
    return image


def trim(image: Image.Image) -> Image.Image:
    alpha = image.getchannel("A")
    if alpha.getextrema()[0] < 255:
        bbox = alpha.getbbox()
    else:
        background = Image.new("RGBA", image.size, image.getpixel((0, 0)))
        difference = ImageChops.difference(image, background).convert("L")
        mask = difference.point(lambda value: 255 if value > 12 else 0)
        bbox = mask.getbbox()
    if bbox:
        image = image.crop(bbox)
    if image.width > 0 and image.height > 0:
        return image
    return Image.new("RGBA", (300, 100), (255, 255, 255, 0))


def normalize(image: Image.Image) -> Image.Image:
    image = trim(image)
    alpha = image.getchannel("A")
    visible = [
        pixel[:3]
        for pixel in image.getdata()
        if pixel[3] > 32
    ]
    mean_luminance = (
        sum(0.2126 * red + 0.7152 * green + 0.0722 * blue for red, green, blue in visible)
        / len(visible)
        if visible
        else 0
    )
    bright_share = (
        sum(1 for red, green, blue in visible if 0.2126 * red + 0.7152 * green + 0.0722 * blue > 230)
        / len(visible)
        if visible
        else 0
    )
    has_transparency = alpha.getextrema()[0] < 255
    needs_dark_surface = has_transparency and (mean_luminance > 205 or bright_share > 0.28)
    background = (10, 22, 40, 255) if needs_dark_surface else (255, 255, 255, 0)
    target = Image.new("RGBA", (600, 220), background)
    scale = min(540 / image.width, 170 / image.height, 1.0 if max(image.size) > 500 else 8.0)
    resized = image.resize(
        (max(1, round(image.width * scale)), max(1, round(image.height * scale))),
        Image.Resampling.LANCZOS,
    )
    target.alpha_composite(resized, ((600 - resized.width) // 2, (220 - resized.height) // 2))
    return target


def official_candidates(domain: str):
    homepage = f"https://{domain}"
    try:
        response = SESSION.get(homepage, timeout=18, allow_redirects=True)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, "html.parser")
        base = response.url
        scored = []
        for image in soup.find_all("img"):
            src = image.get("src") or image.get("data-src") or image.get("data-lazy-src")
            if not src:
                continue
            haystack = " ".join(
                [
                    src,
                    image.get("alt") or "",
                    image.get("id") or "",
                    " ".join(image.get("class") or []),
                ]
            ).lower()
            score = 0
            if "logo" in haystack:
                score += 10
            if "header" in haystack or "brand" in haystack:
                score += 3
            if "footer" in haystack:
                score -= 2
            if score > 0:
                scored.append((score, urljoin(base, src), "official-site-logo"))
        for link in soup.find_all("link"):
            rel = " ".join(link.get("rel") or []).lower()
            href = link.get("href")
            if href and ("icon" in rel or "apple-touch-icon" in rel):
                scored.append((2, urljoin(base, href), "official-site-icon"))
        for _, url, source_type in sorted(scored, reverse=True)[:12]:
            yield url, source_type
    except Exception:
        return


def fallback_candidates(domain: str):
    yield f"https://logo.clearbit.com/{domain}?size=512", "clearbit-domain-logo"
    yield f"https://www.google.com/s2/favicons?domain={domain}&sz=256", "google-domain-icon"


def override_candidates(domain: str):
    for url in VERIFIED_OVERRIDES.get(domain, []):
        yield url, "verified-logo-override"


def fetch_logo(domain: str):
    errors = []
    seen = set()
    clearbit = [(f"https://logo.clearbit.com/{domain}?size=512", "clearbit-domain-logo")]
    google = [(f"https://www.google.com/s2/favicons?domain={domain}&sz=256", "google-domain-icon")]
    candidates = [
        *override_candidates(domain),
        *clearbit,
        *official_candidates(domain),
        *google,
    ]
    for url, source_type in candidates:
        if url in seen:
            continue
        seen.add(url)
        try:
            response = SESSION.get(url, timeout=20, allow_redirects=True)
            response.raise_for_status()
            if len(response.content) < 200:
                raise ValueError("response too small")
            image = image_from_response(response)
            return normalize(image), response.url, source_type
        except Exception as exc:
            errors.append(f"{source_type}: {type(exc).__name__}")
    raise RuntimeError("; ".join(errors[-6:]))


def main():
    wb = load_workbook(SOURCE, data_only=True)
    ws = wb["Top 30 Shortlist"]
    manifest = []
    for row in range(8, 38):
        rank = int(ws.cell(row, 1).value)
        client = str(ws.cell(row, 3).value)
        domain = str(ws.cell(row, 18).value)
        output = OUT / f"{slug(domain)}.png"
        record = {"rank": rank, "client": client, "domain": domain, "file": output.name}
        try:
            image, source_url, source_type = fetch_logo(domain)
            image.save(output, "PNG", optimize=True)
            record.update({"status": "ok", "source_type": source_type, "source_url": source_url})
        except Exception as exc:
            record.update({"status": "failed", "error": str(exc)})
        manifest.append(record)
        print(json.dumps({k: record.get(k) for k in ("rank", "client", "status", "source_type")}))
        time.sleep(0.15)
    (OUT / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    ok = sum(item["status"] == "ok" for item in manifest)
    print(f"Fetched {ok}/30 logos")
    if ok < 30:
        raise SystemExit(f"Logo fetch incomplete: {ok}/30")


if __name__ == "__main__":
    main()
